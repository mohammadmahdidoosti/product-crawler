"""Catalog input, step 1: a supplier's price list (PDF, Excel or CSV) -> tables, text lines and pictures.

Writes C1_read.json in the run folder:
  tables  [{page, title, rows: [[cell, ...]], row_y: [...], bbox: [x0, top, x1, bottom]}]
  lines   [{page, n, y, text}]          every text line, numbered (n) for the language model
  images  [{page, file, y0, y1, repeated}]  pictures inside the catalog (C1_img/), with where they sit
  pages   number of pages; page PNGs in C1_pages/ (for the vision model, scanned pages)
  reversed  True when the PDF stored its Persian text in visual order (PowerPoint/Corel exports:
            'تمیق' for 'قیمت'); every line and cell is then put back in reading order here

Persian numbers: '/' is the decimal point there ('2/5mm' = 2.5 mm, '0/7' = 0.7) while '1/2' in an
inch column is a half; values are only cleaned here, their meaning is settled in catalog.py.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import unicodedata
from pathlib import Path

log = logging.getLogger(__name__)

FA = re.compile(r"[؀-ۿ]")
# words a price list nearly always has; seen backwards ('تمیق') they show the text is in visual order
SIGNALS = ("قیمت", "لیست", "سایز", "مدل", "تعداد", "ضخامت", "قطر", "طول", "ارتفاع", "نوع", "اینچ", "تومان",
           "ریال", "ظرفیت", "شرح", "واحد", "لوله", "ورودی", "عرض", "وزن", "کالا", "محصول", "ابعاد")
# a left-to-right run inside a Persian line: digits, Latin, units and the signs between them
LTR_RUN = re.compile(r"[0-9A-Za-z½¼¾⅓⅔][0-9A-Za-z½¼¾⅓⅔.,/*×%+:'\"_-]*(?:\s+[0-9A-Za-z½¼¾⅓⅔.,/*×%+:'\"_-]+)*")
LETTERS = str.maketrans({"ي": "ی", "ى": "ی", "ك": "ک", "‏": "", "‎": ""})


def clean(s) -> str:
    """Presentation forms (U+FBxx/U+FExx) to plain letters, Arabic ی/ک to Persian."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s)).translate(LETTERS)
    return s.replace("\r", "").strip()


def is_reversed(texts: list[str]) -> bool:
    joined = " ".join(texts)
    ok = sum(joined.count(w) for w in SIGNALS)
    back = sum(joined.count(w[::-1]) for w in SIGNALS)
    return back > ok


def unreverse_line(line: str) -> str:
    """A Persian line stored in visual order -> reading order; Latin and number runs keep theirs."""
    if not FA.search(line):
        return line
    r = line[::-1]
    # brackets need no mirroring: the glyph stored is already the one for reading order
    return LTR_RUN.sub(lambda m: m.group(0)[::-1], r)


def fix(s, rev: bool) -> str:
    s = clean(s)
    if rev:
        s = "\n".join(unreverse_line(x) for x in s.split("\n"))
    return re.sub(r"[ \t]+", " ", s).strip()


# ------------------------------------------------------------------ PDF

def read_pdf(path: Path, out: Path, render_dpi: int = 110) -> dict:
    import pdfplumber
    pages, raw_lines, raw_tables, raw_images = 0, [], [], []
    (out / "C1_pages").mkdir(parents=True, exist_ok=True)
    (out / "C1_img").mkdir(parents=True, exist_ok=True)
    with pdfplumber.open(str(path)) as pdf:
        pages = len(pdf.pages)
        # (from the raw glyphs: extract_text() would cache a text map built before the fix below)
        rev = is_reversed([clean("".join(c["text"] for c in p.chars)) for p in pdf.pages])
        for i, p in enumerate(pdf.pages, 1):
            W, H = float(p.width), float(p.height)
            if rev:
                # a ligature glyph ('تی', 'لا', 'سی') carries its letters in reading order already: turned
                # here, so reversing the whole line later puts it back the right way round
                for c in p.chars:
                    if len(c["text"]) > 1 and FA.search(c["text"]):
                        c["text"] = c["text"][::-1]
            tables = []
            for t in p.find_tables():
                rows = t.extract()
                if len(rows) < 2 or max(len(r) for r in rows) < 2:
                    continue
                tables.append({"page": i, "bbox": [round(float(x), 1) for x in t.bbox], "rows": rows,
                               "row_y": [round(float(r.bbox[1]), 1) for r in t.rows]})
            raw_tables += tables
            for ln in p.extract_text_lines():
                y = float(ln["top"])
                inside = any(tb["bbox"][1] - 1 <= y <= tb["bbox"][3] and tb["bbox"][0] - 1 <= float(ln["x0"]) <= tb["bbox"][2]
                             for tb in tables)
                raw_lines.append({"page": i, "y": round(y, 1), "text": ln["text"], "in_table": inside})
            for k, im in enumerate(p.images):
                x0, top, x1, bottom = (max(0.0, float(im["x0"])), max(0.0, float(im["top"])),
                                       min(W, float(im["x1"])), min(H, float(im["bottom"])))
                w, h = x1 - x0, bottom - top
                if w < 40 or h < 40 or w * h > 0.6 * W * H:          # icons, bullets, full-page backgrounds
                    continue
                try:
                    data = im["stream"].get_data()
                except Exception:                                    # noqa: BLE001
                    data = b""
                digest = hashlib.md5(data or f"{i}-{k}".encode()).hexdigest()[:12]
                raw_images.append({"page": i, "k": k, "box": (x0, top, x1, bottom), "hash": digest})
            try:
                p.to_image(resolution=render_dpi).save(str(out / "C1_pages" / f"{i:03d}.png"))
            except Exception as e:                                   # noqa: BLE001 - rendering is a help only
                log.info("page %d not rendered: %s", i, e)
        # a picture on many pages is the catalog's own decoration (logo, background, contacts)
        count: dict[str, set] = {}
        for x in raw_images:
            count.setdefault(x["hash"], set()).add(x["page"])
        many = max(3, int(0.3 * pages))
        images = []
        for x in raw_images:
            rep = len(count[x["hash"]]) >= many
            f = ""
            if not rep:
                try:
                    pic = pdf.pages[x["page"] - 1].crop(x["box"]).to_image(resolution=200).original
                    f = f"C1_img/p{x['page']:03d}_{x['k']}.png"
                    pic.save(str(out / f))
                except Exception as e:                               # noqa: BLE001
                    log.info("catalog picture p%d/%d not saved: %s", x["page"], x["k"], e)
            images.append({"page": x["page"], "file": f, "y0": round(x["box"][1], 1), "y1": round(x["box"][3], 1),
                           "repeated": rep})
    lines = [{"page": x["page"], "y": x["y"], "text": fix(x["text"], rev), "in_table": x["in_table"]} for x in raw_lines]
    for t in raw_tables:
        t["rows"] = [[None if c is None else fix(c, rev) for c in r] for r in t["rows"]]
        t["title"] = _title(lines, t)
    return _finish({"kind": "pdf", "pages": pages, "reversed": rev, "tables": raw_tables, "lines": lines,
                    "images": images})


def _title(lines: list[dict], t: dict) -> str:
    """The heading of a table: the nearest line above it on its page ('لیست قیمت …' preferred)."""
    above = [x for x in lines if x["page"] == t["page"] and x["y"] < t["bbox"][1] and not x["in_table"]
             and FA.search(x["text"])]
    named = [x for x in above if "قیمت" in x["text"] or "لیست" in x["text"]]
    pick = (named or above)
    if pick:
        return max(pick, key=lambda x: x["y"])["text"]
    # a table drawn up to the page's top: the price-list heading nearest to it, wherever it is
    near = [x for x in lines if x["page"] == t["page"] and "لیست قیمت" in x["text"]]
    return min(near, key=lambda x: abs(x["y"] - t["bbox"][1]))["text"] if near else ""


# ------------------------------------------------------------------ Excel / CSV

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return clean(v)


def read_excel(path: Path, out: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    tables, lines, images = [], [], []
    (out / "C1_img").mkdir(parents=True, exist_ok=True)
    for si, ws in enumerate(wb.worksheets, 1):
        merged = {}
        for rg in ws.merged_cells.ranges:                       # a merged cell's value fills its whole area
            v = ws.cell(rg.min_row, rg.min_col).value
            for r in range(rg.min_row, rg.max_row + 1):
                for c in range(rg.min_col, rg.max_col + 1):
                    merged[(r, c)] = v
        rows, row_y = [], []
        for r_i, row in enumerate(ws.iter_rows(), 1):
            cells = [_cell(merged.get((r_i, c.column), c.value)) for c in row]
            if not any(cells):
                continue
            rows.append(cells)
            row_y.append(float(r_i))
            if sum(1 for c in cells if c) == 1:
                lines.append({"page": si, "y": float(r_i), "text": next(c for c in cells if c), "in_table": False})
        title = ""
        while rows and sum(1 for c in rows[0] if c) == 1:      # heading rows above the table: its title
            one = next(c for c in rows[0] if c)
            if not title or "قیمت" in one or "لیست" in one:
                title = one
            rows.pop(0)
            row_y.pop(0)
        if rows:
            width = max(len(r) for r in rows)
            used = [j for j in range(width) if any(j < len(r) and r[j] for r in rows)]
            rows = [[r[j] if j < len(r) else "" for j in used] for r in rows]
            tables.append({"page": si, "bbox": [0, row_y[0], 1, row_y[-1]], "rows": rows, "row_y": row_y,
                           "title": title or ws.title, "sheet": ws.title})
        for k, img in enumerate(getattr(ws, "_images", []) or []):
            try:
                data = img._data()
                row = float(img.anchor._from.row + 1)
                f = f"C1_img/s{si:02d}_{k}.png"
                from PIL import Image
                Image.open(io.BytesIO(data)).save(str(out / f))
                images.append({"page": si, "file": f, "y0": row, "y1": row + 1, "repeated": False})
            except Exception as e:                                # noqa: BLE001
                log.info("excel picture %d/%d not saved: %s", si, k, e)
    return _finish({"kind": "excel", "pages": len(wb.worksheets), "reversed": False, "tables": tables,
                    "lines": lines, "images": images})


def read_csv_file(path: Path) -> dict:
    import pandas as pd
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1256", "utf-16"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("رمزگذاری فایل CSV شناخته نشد؛ از Excel با «CSV UTF-8» ذخیره کنید")
    first = text.split("\n", 1)[0]
    sep = max((",", ";", "\t"), key=first.count)
    df = pd.read_csv(io.StringIO(text), sep=sep, dtype=str, keep_default_na=False, header=None)
    rows = [[clean(c) for c in r] for r in df.values.tolist() if any(str(c).strip() for c in r)]
    return _finish({"kind": "csv", "pages": 1, "reversed": False, "lines": [], "images": [],
                    "tables": [{"page": 1, "bbox": [0, 1, 1, len(rows)], "rows": rows,
                                "row_y": [float(i) for i in range(1, len(rows) + 1)], "title": path.stem}]})


def _finish(d: dict) -> dict:
    for n, x in enumerate(d["lines"], 1):
        x["n"] = n
    return d


def read(path: Path, out: Path) -> dict:
    """Any supported file -> the C1 structure (also saved as C1_read.json)."""
    path, out = Path(path), Path(out)
    out.mkdir(parents=True, exist_ok=True)
    head = path.read_bytes()[:8]
    suf = path.suffix.lower()
    if head.startswith(b"%PDF") or suf == ".pdf":
        d = read_pdf(path, out)
    elif head[:2] == b"PK" or suf in (".xlsx", ".xlsm"):
        d = read_excel(path, out)
    elif suf == ".xls":
        raise ValueError("فایل .xls قدیمی است؛ در Excel با Save As نوع xlsx ذخیره کنید")
    else:
        d = read_csv_file(path)
    d["source"] = path.name
    (out / "C1_read.json").write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
    log.info("catalog read: %s — %d page(s), %d table(s), %d line(s), %d own picture(s)%s", path.name, d["pages"],
             len(d["tables"]), len(d["lines"]), sum(1 for x in d["images"] if x["file"]),
             " (Persian text was in visual order: fixed)" if d["reversed"] else "")
    return d
