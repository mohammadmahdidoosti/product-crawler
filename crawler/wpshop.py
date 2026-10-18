# Vendored unchanged from the main branch of this repo (the team's WordPress product crawler,
# product-crawler-vutt.onrender.com). Used here for its product-page parser and URL rules.
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Product crawler for WordPress / WooCommerce (and generic) shop sites.

For every site it tries, in order:
  1. WooCommerce Store API  (/wp-json/wc/store/v1/products)  -> fastest, cleanest data
  2. XML sitemaps            (product-sitemap.xml, wp-sitemap-posts-product-*.xml, ...)
  3. Link crawling (BFS)     starting from the home page / shop page
Product pages are parsed from JSON-LD (schema.org Product), OpenGraph and
WooCommerce HTML markup, so it also works on non-WooCommerce sites that expose
structured data.

Output (in ./output by default):
  <domain>.jsonl      one product per line, written as it is crawled (resumable)
  <domain>.csv        per-site table (UTF-8 with BOM, opens correctly in Excel)
  all_products.csv    every site combined
  all_products.xlsx   every site combined (if openpyxl is installed)

Usage:
  python crawler.py                         # crawl all default sites
  python crawler.py --sites https://tsaco-diesel.com/ https://parsshoa.co/
  python crawler.py --mode sitemap --delay 2 --max-pages 3000 --resume
  python crawler.py --proxy http://127.0.0.1:8080
"""

from __future__ import annotations

import argparse
import csv
import gzip
import html as html_lib
import json
import logging
import os
import random
import re
import sys
import threading
import time
import xml.etree.ElementTree as ET
from collections import deque
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, as_completed, wait
from datetime import datetime, timezone
from urllib.parse import parse_qsl, unquote, urldefrag, urlencode, urljoin, urlparse

import lxml.html
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

DEFAULT_SITES = [
    "https://tsaco-diesel.com/",
    "https://parsshoa.co/",
    "https://foolad-parsekaveh.com/",
    "https://mahyaelectronic.ir/",
    "https://yahyabrand.com/",
    "https://alvandlift.ir/",
    "https://safetypatan.com/",
    "https://saterco.ir/",
]

DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/129.0 Safari/537.36 ProductCrawler/1.0"
)
ROBOTS_AGENT_TOKEN = "ProductCrawler"

FIELDS = [
    "site", "source", "url", "product_id", "name", "sku", "price",
    "regular_price", "sale_price", "min_price", "max_price", "currency",
    "price_text", "in_stock", "categories", "tags", "brand", "short_description",
    "description", "description_html", "specs", "image_main", "images", "image_count",
    "description_images", "files", "meta_title", "meta_description", "image_files",
    "crawled_at",
]
MAX_IMAGES_PER_PRODUCT = 15
FILE_EXT_RE = re.compile(r"\.(pdf|docx?|xlsx?|pptx?|zip|rar|dwg|dxf|step|stp|igs)(?:$|\?)", re.I)
PLACEHOLDER_IMG_RE = re.compile(r"lazy|placeholder|blank\.(gif|png)|spinner|loading|data:image|1x1|pixel|"
                                r"woocommerce-placeholder|/emoji/|gravatar", re.I)

# URL path fragments that usually indicate a single product page
PRODUCT_URL_PATTERNS = [
    r"/product/", r"/products/[^/]+", r"/shop/[^/]+/?$", r"/%D9%85%D8%AD%D8%B5%D9%88%D9%84/",
    r"/محصول/", r"/item/", r"/p/\d+", r"/kala/", r"/کالا/",
]
# URL path fragments that are never products and not worth crawling
EXCLUDE_URL_PATTERNS = [
    r"/wp-admin", r"/wp-login", r"/wp-json", r"/wp-content/", r"/wp-includes/",
    r"/cart", r"/checkout", r"/my-account", r"/account", r"/login", r"/register",
    r"/feed", r"/comments/", r"/author/", r"/xmlrpc", r"/wishlist", r"/compare",
    r"/سبد-خرید", r"/تسویه-حساب", r"/حساب-کاربری",
    r"add-to-cart=", r"add_to_wishlist", r"replytocom=", r"orderby=", r"filter_",
    r"[?&]s=", r"[?&]min_price=", r"[?&]max_price=", r"/tag/", r"/product-tag/",
    r"\.(?:jpe?g|png|gif|webp|svg|pdf|zip|rar|mp4|mp3|avi|docx?|xlsx?|css|js|ico|woff2?|ttf)(?:$|\?)",
]
# Query keys kept when normalising URLs (pagination); everything else is dropped
KEEP_QUERY_KEYS = {"page", "paged", "product-page", "p", "pg"}

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

log = logging.getLogger("crawler")


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #

def to_latin_digits(s: str) -> str:
    return (s or "").translate(PERSIAN_DIGITS)


BLOCK_TAGS = {"p", "div", "li", "ul", "ol", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "table", "dt", "dd",
              "blockquote", "section", "article", "figure", "figcaption", "hr", "pre"}


def html_to_text(el) -> str:
    """Readable text: newlines only between block elements, ' | ' between table cells."""
    if el is None:
        return ""
    el = BeautifulSoup(str(el), "html.parser")
    for bad in el.find_all(["script", "style", "noscript"]):
        bad.decompose()
    for br in el.find_all("br"):
        br.replace_with("\n")
    for cell in el.find_all(["td", "th"]):
        cell.insert_after(" | ")
    for blk in el.find_all(BLOCK_TAGS):
        blk.insert_before("\n")
        blk.insert_after("\n")
    text = el.get_text("")
    text = re.sub(r"[ \t]*\|[ \t]*\n", "\n", text)
    return text


def clean_text(s) -> str:
    if s is None:
        return ""
    if not isinstance(s, str):
        if hasattr(s, "find_all"):          # a BeautifulSoup element
            s = html_to_text(s)
        else:
            s = str(s)
    elif "<" in s and ">" in s:
        s = html_to_text(s)
    s = html_lib.unescape(s).replace("\xa0", " ")
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n", s)
    return s.strip()


def parse_number(token: str):
    """'1,800,000' / '۱٬۸۰۰٬۰۰۰' / '11.000.000.000' / '12.5' -> float"""
    t = to_latin_digits(token).strip()
    t = t.replace("٬", ",").replace("،", ",").replace(" ", "")
    if not t:
        return None
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", t):           # dots as thousands sep
        t = t.replace(".", "")
    t = t.replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def extract_numbers(text: str) -> list[float]:
    t = to_latin_digits(text or "")
    out = []
    for tok in re.findall(r"\d[\d,٬،.]*\d|\d", t):
        n = parse_number(tok)
        if n is not None:
            out.append(n)
    return out


def detect_currency(text: str) -> str:
    t = text or ""
    if "تومان" in t or "toman" in t.lower():
        return "IRT"
    if "ریال" in t or "﷼" in t or "rial" in t.lower():
        return "IRR"
    return ""


def norm_host(host: str) -> str:
    host = (host or "").lower()
    return host[4:] if host.startswith("www.") else host


_SECOND_LEVEL = {"co", "ac", "org", "gov", "net", "com", "sch", "id", "edu"}
SKIP_SUBDOMAINS = ("blog.", "forum.", "mail.", "webmail.", "cdn.", "static.", "status.", "help.", "support.",
                   "cpanel.", "ftp.", "api.", "admin.", "panel.")


def base_domain(host: str) -> str:
    """'list.invt-co.com' -> 'invt-co.com', 'shop.example.co.ir' -> 'example.co.ir'."""
    parts = norm_host(host).split(":")[0].split(".")
    if len(parts) >= 3 and parts[-2] in _SECOND_LEVEL and len(parts[-1]) == 2:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def normalize_url(url: str) -> str:
    url, _ = urldefrag(url)
    p = urlparse(url)
    q = [(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True) if k in KEEP_QUERY_KEYS]
    path = p.path or "/"
    return p._replace(path=path, query=urlencode(q), fragment="").geturl()


def is_excluded(url: str) -> bool:
    u = unquote(url).lower()
    return any(re.search(pat, u) for pat in EXCLUDE_URL_PATTERNS)


def looks_like_product_url(url: str) -> bool:
    path = urlparse(url).path
    if re.search(r"/product-(?:cat|category|tag|brand)/", path, re.I):
        return False
    return any(re.search(p, path, re.I) or re.search(p, unquote(path), re.I) for p in PRODUCT_URL_PATTERNS)


def looks_like_listing_url(url: str) -> bool:
    u = unquote(url).lower()
    return bool(re.search(r"/(shop|store|products?|product-cat|product-category|category|"
                          r"فروشگاه|محصولات|دسته)|/page/\d+|[?&](page|paged|product-page)=", u))


# --------------------------------------------------------------------------- #
# robots.txt (supports * and $ wildcards, longest-match rule)
# --------------------------------------------------------------------------- #

class Robots:
    def __init__(self, text: str = "", agent: str = ROBOTS_AGENT_TOKEN):
        self.rules: list[tuple[bool, str]] = []
        self.crawl_delay: float | None = None
        self.sitemaps: list[str] = []
        self._parse(text or "", agent.lower())

    def _parse(self, text: str, agent: str):
        groups = []           # (agents, rules, delay)
        agents, rules, delay = [], [], None
        last_was_agent = False
        for raw in text.splitlines():
            line = raw.split("#", 1)[0].strip()
            if ":" not in line:
                continue
            key, val = line.split(":", 1)
            key, val = key.strip().lower(), val.strip()
            if key == "user-agent":
                if not last_was_agent and agents:
                    groups.append((agents, rules, delay))
                    agents, rules, delay = [], [], None
                agents.append(val.lower())
                last_was_agent = True
                continue
            last_was_agent = False
            if key in ("allow", "disallow"):
                rules.append((key == "allow", val))
            elif key == "crawl-delay":
                try:
                    delay = float(val)
                except ValueError:
                    pass
            elif key == "sitemap":
                self.sitemaps.append(val)
        if agents:
            groups.append((agents, rules, delay))

        specific = [g for g in groups if any(a != "*" and a in agent for a in g[0])]
        chosen = specific or [g for g in groups if "*" in g[0]]
        for _, r, d in chosen:
            self.rules.extend(r)
            if d is not None:
                self.crawl_delay = d

    @staticmethod
    def _match_len(pattern: str, path: str) -> int:
        pattern = unquote(pattern)
        anchored = pattern.endswith("$")
        if anchored:
            pattern = pattern[:-1]
        regex = "^" + "".join(".*" if c == "*" else re.escape(c) for c in pattern)
        if anchored:
            regex += "$"
        return len(pattern) if re.match(regex, path) else -1

    def allowed(self, url: str) -> bool:
        p = urlparse(url)
        path = unquote(p.path or "/") + (("?" + unquote(p.query)) if p.query else "")
        best_len, best_allow = -1, True
        for allow, pat in self.rules:
            if not pat:
                continue
            n = self._match_len(pat, path)
            if n > best_len or (n == best_len and allow):
                best_len, best_allow = n, allow
        return best_allow if best_len >= 0 else True


# --------------------------------------------------------------------------- #
# HTTP fetching with retries + per-site rate limit
# --------------------------------------------------------------------------- #

class Fetcher:
    def __init__(self, args, stop_event: threading.Event | None = None):
        self.stop_event = stop_event
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": args.user_agent,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
            "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.7",
        })
        retry = Retry(total=3, connect=3, read=3, backoff_factor=1.5,
                      status_forcelist=(429, 500, 502, 503, 504),
                      allowed_methods=("GET", "HEAD"), respect_retry_after_header=True)
        adapter = HTTPAdapter(max_retries=retry, pool_connections=10, pool_maxsize=10)
        self.session.mount("http://", adapter)
        self.session.mount("https://", adapter)
        if args.proxy:
            self.session.proxies.update({"http": args.proxy, "https": args.proxy})
        self.session.verify = not args.insecure
        self.timeout = args.timeout
        self.delay = args.delay
        self._lock = threading.Lock()
        self._next = 0.0

    def get(self, url: str):
        if self.stop_event is not None and self.stop_event.is_set():
            return None
        with self._lock:                          # polite spacing between requests
            wait = self._next - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            self._next = time.monotonic() + self.delay * random.uniform(0.8, 1.3)
        try:
            return self.session.get(url, timeout=self.timeout, allow_redirects=True)
        except requests.RequestException as e:
            log.debug("request failed %s: %s", url, e)
            return None


# --------------------------------------------------------------------------- #
# Product page parsing
# --------------------------------------------------------------------------- #

def _iter_jsonld(soup: BeautifulSoup):
    for tag in soup.find_all("script", attrs={"type": re.compile(r"ld\+json", re.I)}):
        raw = tag.string or tag.get_text() or ""
        raw = re.sub(r"[\x00-\x1f]", " ", raw).strip()
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except ValueError:
            continue
        stack = [data]
        while stack:
            node = stack.pop()
            if isinstance(node, list):
                stack.extend(node)
            elif isinstance(node, dict):
                yield node
                for k in ("@graph", "mainEntity", "itemListElement", "item"):
                    if k in node:
                        stack.append(node[k])


def _is_type(node: dict, name: str) -> bool:
    t = node.get("@type")
    types = t if isinstance(t, list) else [t]
    return any(isinstance(x, str) and x.lower() == name.lower() for x in types)


def find_jsonld_product(soup: BeautifulSoup) -> dict | None:
    best = None
    for node in _iter_jsonld(soup):
        if _is_type(node, "Product") or _is_type(node, "ProductGroup"):
            if best is None or (node.get("offers") and not best.get("offers")):
                best = node
    return best


def _as_list(x):
    if x is None:
        return []
    return x if isinstance(x, list) else [x]


def _image_urls(x) -> list[str]:
    out = []
    for item in _as_list(x):
        if isinstance(item, str):
            out.append(item)
        elif isinstance(item, dict):
            u = item.get("url") or item.get("contentUrl") or item.get("@id")
            if isinstance(u, str):
                out.append(u)
    return out


def _dedupe(seq):
    seen, out = set(), []
    for x in seq:
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    return out


def first(soup, selectors: list[str]):
    """First element matching the selectors, tried in priority order."""
    for sel in selectors:
        el = soup.select_one(sel)
        if el is not None:
            return el
    return None


def best_img_src(img) -> str:
    """Real image URL of an <img>, looking past lazy-load placeholders and picking the largest srcset entry."""
    for attr in ("data-large_image", "data-zoom-image", "data-full", "data-src", "data-lazy-src",
                 "data-original", "src"):
        u = (img.get(attr) or "").strip()
        if u and not PLACEHOLDER_IMG_RE.search(u):
            return u
    for attr in ("data-srcset", "data-lazy-srcset", "srcset"):
        srcset = img.get(attr) or ""
        best, best_w = "", -1
        for part in srcset.split(","):
            bits = part.strip().split()
            if not bits:
                continue
            w = int(re.sub(r"\D", "", bits[1]) or 0) if len(bits) > 1 else 0
            if w > best_w and not PLACEHOLDER_IMG_RE.search(bits[0]):
                best, best_w = bits[0], w
        if best:
            return best
    return ""


def strip_wp_size(u: str) -> str:
    """'photo-300x300.jpg' -> 'photo.jpg' (WordPress thumbnails -> original upload)."""
    return re.sub(r"-\d{2,4}x\d{2,4}(?=\.(?:jpe?g|png|webp|gif|avif)(?:$|\?))", "", u, flags=re.I)


SAFE_TAGS = {"p", "br", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "strong", "b", "em", "i",
             "u", "a", "img", "table", "thead", "tbody", "tfoot", "tr", "th", "td", "blockquote",
             "span", "div", "figure", "figcaption", "hr", "dl", "dt", "dd", "sup", "sub", "video", "source"}
SAFE_ATTRS = {"href", "src", "alt", "title", "colspan", "rowspan"}


def sanitize_html(el, base_url: str) -> str:
    """Clean copy of an element's inner HTML: no scripts/styles/classes, absolute links, real image URLs."""
    if el is None:
        return ""
    root = BeautifulSoup(str(el), "html.parser").find(True)   # detached copy of el
    if root is None:
        return ""
    for bad in root.find_all(["script", "style", "noscript", "iframe", "form", "button", "svg",
                              "input", "select", "textarea", "link", "meta"]):
        bad.decompose()
    for img in root.find_all("img"):
        src = best_img_src(img)
        if src:
            img["src"] = urljoin(base_url, src)
        else:
            img.decompose()
    for tag in root.find_all(True):
        if tag.name not in SAFE_TAGS:
            tag.unwrap()
            continue
        tag.attrs = {k: v for k, v in tag.attrs.items() if k in SAFE_ATTRS}
        for k in ("href", "src"):
            if tag.get(k):
                tag[k] = urljoin(base_url, tag[k])
    html = re.sub(r"\n\s*\n+", "\n", root.decode_contents())
    return re.sub(r"<(p|div|span)>\s*</\1>", "", html).strip()


def content_images(el, base_url: str) -> list[str]:
    if el is None:
        return []
    out = []
    for img in el.find_all("img"):
        u = best_img_src(img)
        if u and not u.lower().endswith(".svg"):
            out.append(urljoin(base_url, u))
    return _dedupe(out)


def extract_specs(*containers) -> dict:
    """Key/value pairs from 2-column tables, <dl> lists and 'key: value' list items."""
    specs: dict[str, str] = {}

    def add(k, v):
        k, v = clean_text(k).strip(" :：-"), clean_text(v)
        if k and v and len(k) <= 60 and k != v and k not in specs:
            specs[k] = v[:1000]

    for c in containers:
        if c is None:
            continue
        for table in c.find_all("table"):
            for tr in table.find_all("tr"):
                cells = tr.find_all(["th", "td"], recursive=False)
                if len(cells) == 2:
                    add(cells[0].get_text(" "), cells[1].get_text(" "))
        for dl in c.find_all("dl"):
            for dt in dl.find_all("dt"):
                dd = dt.find_next_sibling("dd")
                if dd:
                    add(dt.get_text(" "), dd.get_text(" "))
        for lst in c.find_all(["ul", "ol"]):
            pairs = []
            for li in lst.find_all("li", recursive=False):
                kids = [k for k in li.find_all(True, recursive=False) if clean_text(k.get_text(" "))]
                if len(kids) == 2 and not kids[0].find(["ul", "ol", "table"]):
                    pairs.append((kids[0].get_text(" "), kids[1].get_text(" ")))
                    continue
                t = clean_text(li.get_text(" "))
                m = re.match(r"^([^:：]{2,40})\s*[:：]\s*(.{1,300})$", t, re.S)
                if m:
                    pairs.append(m.groups())
            if len(pairs) >= 2:            # a single "x: y" bullet is usually prose, not a spec list
                for k, v in pairs:
                    add(k, v)
    return specs


def find_files(soup, base_url: str) -> list[str]:
    scope = soup.select_one("main, #main, .site-main, .product, article") or soup
    out = []
    for a in scope.find_all("a", href=True):
        if FILE_EXT_RE.search(a["href"]):
            out.append(urljoin(base_url, a["href"]))
    return _dedupe(out)


def set_images(rec: dict, urls: list[str]):
    urls = _dedupe(u for u in urls if u and not PLACEHOLDER_IMG_RE.search(u))
    # drop WordPress thumbnails when the original upload is also in the list
    originals = set(urls)
    urls = [u for u in urls if strip_wp_size(u) == u or strip_wp_size(u) not in originals]
    rec["images"] = " | ".join(urls)
    rec["image_main"] = urls[0] if urls else ""
    rec["image_count"] = len(urls)


SKU_LABEL_RE = re.compile(r"^(کد محصول|کد کالا|شناسه محصول|شناسه کالا|کد|sku|product code|part number|پارت نامبر)$", re.I)
BRAND_LABEL_RE = re.compile(r"^(برند|سازنده|نام تجاری|brand|manufacturer)$", re.I)
GENERIC_CRUMB_RE = re.compile(r"^(خانه|صفحه اصلی|home|فروشگاه( .{1,30})?|shop|محصولات|products|دسته بندی محصولات|"
                              r"دسته‌بندی محصولات|همه محصولات|لیست کالا ?ها)$", re.I)


def labeled_values(soup) -> dict:
    """'label : value' snippets in small elements, e.g. <div>کد محصول : <b>100282</b></div>."""
    out = {}
    for el in soup.find_all(["div", "li", "p", "span", "td"]):
        if el.find(["div", "li", "p", "table", "ul"]):
            continue
        t = clean_text(el.get_text(" "))
        m = re.match(r"^([^:：]{2,30})\s*[:：]\s*(.{1,120})$", t)
        if m and m.group(1).strip() not in out:
            out[m.group(1).strip()] = m.group(2).strip()
    return out


_LXML_PARSER = lxml.html.HTMLParser(encoding="utf-8", remove_comments=True)
_JUNK_CLASS_RE = re.compile(r"mega-?menu|mobile-?(menu|nav)|menu-?mobile|off-?canvas|side-?menu|navbar-collapse|"
                            r"cart-?(widget|sidebar)|mini-?cart|search-?(modal|popup)", re.I)


def preclean(html: str) -> str:
    """Strip what the extractor never reads (scripts, styles, SVG icons, site menus, header/footer) with lxml,
    which is ~15x faster than building a BeautifulSoup tree. Iranian shop themes often carry thousands of
    menu nodes per page; dropping them first makes the detailed parse several times cheaper."""
    if not html or len(html) < 20000:
        return html
    try:
        root = lxml.html.fromstring(html.encode("utf-8", "replace"), parser=_LXML_PARSER)
    except Exception:
        return html
    if root is None or root.tag != "html":
        return html
    drop = []
    for el in root.iter("script", "style", "noscript", "svg", "iframe", "template", "link"):
        if el.tag == "script" and "json" in (el.get("type") or "").lower():
            continue                          # JSON-LD and embedded JSON data are needed
        drop.append(el)
    for el in root.iter("header", "footer", "nav", "aside", "div", "ul", "section"):
        if el.tag in ("div", "ul", "section"):
            if not _JUNK_CLASS_RE.search(f"{el.get('class', '')} {el.get('id', '')}"):
                continue
        elif el.tag == "nav" and re.search(r"breadcrumb", f"{el.get('aria-label', '')} {el.get('class', '')}", re.I):
            continue
        if any(a.tag in ("main", "article") for a in el.iterancestors()):
            continue
        if el.find(".//h1") is not None or el.xpath(".//*[contains(@class,'product_title') or contains(@class,'breadcrumb')]"):
            continue
        drop.append(el)
    for el in drop:
        if el.getparent() is not None:
            el.drop_tree()          # keeps the element's tail text
    try:
        return lxml.html.tostring(root, encoding="unicode", method="html")
    except Exception:
        return html


def parse_product_page_bs4(html: str, url: str, site: str, trust_as_product: bool = False) -> dict | None:
    """Previous BeautifulSoup implementation, kept for comparison tests."""
    soup = BeautifulSoup(preclean(html), "lxml")
    ld = find_jsonld_product(soup)
    meta = {}
    for m in soup.find_all("meta"):
        k = (m.get("property") or m.get("name") or m.get("itemprop") or "").strip().lower()
        if k and m.get("content") and k not in meta:
            meta[k] = m["content"].strip()
    body_cls = " ".join((soup.body.get("class") or [])) if soup.body else ""

    html_signals = bool(
        soup.select_one("form.cart, .product_title, .single-product-summary, "
                        "[itemtype*='schema.org/Product'], .woocommerce-product-gallery")
    )
    is_product = (
        ld is not None
        or meta.get("og:type", "").lower() == "product"
        or "single-product" in body_cls
        or html_signals
    )
    generic = not is_product and (trust_as_product or looks_like_product_url(url))
    if not is_product and not generic:
        return None
    for junk in soup.select("header, footer, nav:not([aria-label*=readcrumb]):not([class*=breadcrumb]), .mega-menu, #mega-menu"):
        if junk.find_parent(["main", "article"]) is None and not junk.select_one("h1"):
            junk.decompose()      # menus repeat on every page; keep them out of names/specs/images

    rec = {k: "" for k in FIELDS}
    rec.update(site=site, source="page", url=url,
               crawled_at=datetime.now(timezone.utc).isoformat(timespec="seconds"))

    # ---- name
    h1 = first(soup, ["h1.product_title", "h1.product-title", ".product_title", "main h1", "h1"])
    crumbs_last = soup.select("nav[aria-label*=readcrumb] li:last-child, [class*=breadcrumb] li:last-child")
    page_title = soup.find("title")
    title_name = ""
    if page_title:
        parts = re.split(r"\s+[-|–—]\s+", clean_text(page_title.get_text()))
        title_name = parts[0] if parts else ""
    rec["name"] = clean_text((ld or {}).get("name")) or (clean_text(h1.get_text()) if h1 else "") \
        or clean_text(meta.get("og:title", "")) \
        or (clean_text(crumbs_last[0].get_text()) if crumbs_last else "") or title_name
    if not rec["name"]:
        return None
    labeled = labeled_values(soup)

    # ---- id / sku
    rec["sku"] = clean_text((ld or {}).get("sku") or (ld or {}).get("mpn") or "")
    if not rec["sku"]:
        el = soup.select_one(".sku, [itemprop=sku]")
        if el:
            rec["sku"] = clean_text(el.get("content") or el.get_text())
        else:
            rec["sku"] = next((v for k, v in labeled.items() if SKU_LABEL_RE.search(k)), "")
    m = re.search(r"postid-(\d+)", body_cls)
    if m:
        rec["product_id"] = m.group(1)
    else:
        el = soup.select_one("[data-product_id], input[name=add-to-cart], button[name=add-to-cart]")
        if el:
            rec["product_id"] = el.get("data-product_id") or el.get("value") or ""

    # ---- brand
    brand = (ld or {}).get("brand")
    if isinstance(brand, list):
        brand = brand[0] if brand else ""
    if isinstance(brand, dict):
        brand = brand.get("name", "")
    rec["brand"] = clean_text(brand) or next((v for k, v in labeled.items() if BRAND_LABEL_RE.search(k)), "")

    # ---- price (JSON-LD offers first)
    prices, currency, availability = [], "", ""
    for off in _as_list((ld or {}).get("offers")):
        if not isinstance(off, dict):
            continue
        currency = currency or off.get("priceCurrency", "")
        availability = availability or str(off.get("availability", ""))
        for key in ("price", "lowPrice", "highPrice"):
            v = off.get(key)
            if v not in (None, ""):
                n = parse_number(str(v))
                if n is not None:
                    prices.append(n)
        spec = off.get("priceSpecification")
        for ps in _as_list(spec):
            if isinstance(ps, dict) and ps.get("price") not in (None, ""):
                n = parse_number(str(ps["price"]))
                if n is not None:
                    prices.append(n)
                currency = currency or ps.get("priceCurrency", "")

    price_el = first(soup, [".summary .price", ".entry-summary .price", ".product-info .price",
                            "p.price", "[itemprop=price]", ".product-price", "span.price"])
    if price_el is None and generic:
        price_el = next((el for el in soup.select("main [class*=price], [class*=price], [id*=price]")
                         if extract_numbers(el.get_text(" ")) and len(el.get_text(" ")) < 120), None)
    price_text = ""
    if price_el:
        price_text = clean_text(price_el.get("content") or price_el.get_text(" "))
        ins, dels = price_el.select_one("ins"), price_el.select_one("del")
        if ins and dels:
            sale = extract_numbers(ins.get_text(" "))
            reg = extract_numbers(dels.get_text(" "))
            if sale:
                rec["sale_price"] = sale[0]
            if reg:
                rec["regular_price"] = reg[0]
    if not price_text and meta.get("product:price:amount"):
        price_text = meta["product:price:amount"]
    rec["price_text"] = price_text

    if not prices and price_text:
        nums = [n for n in extract_numbers(price_text) if n > 0]
        # with sale markup the "current" price is the <ins> one
        prices = [rec["sale_price"]] if rec["sale_price"] != "" else nums
    if prices:
        rec["min_price"], rec["max_price"] = min(prices), max(prices)
        rec["price"] = rec["sale_price"] if rec["sale_price"] != "" else min(prices)
    rec["currency"] = currency or meta.get("product:price:currency", "") or detect_currency(price_text)

    # ---- stock
    av = availability.lower()
    stock_el = soup.select_one(".summary .stock, p.stock, .stock")
    stock_txt = clean_text(stock_el.get_text()) if stock_el else ""
    stock_cls = " ".join(stock_el.get("class") or []) if stock_el else ""
    if "outofstock" in av or "out-of-stock" in stock_cls or "ناموجود" in stock_txt:
        rec["in_stock"] = False
    elif "instock" in av or "in-stock" in stock_cls or "موجود" in stock_txt:
        rec["in_stock"] = True

    # ---- categories
    cats = [clean_text(a.get_text()) for a in soup.select(".product_meta .posted_in a, .posted_in a")]
    if not cats:
        crumbs = soup.select(".woocommerce-breadcrumb a, nav.rank-math-breadcrumb a, "
                             ".breadcrumb a, .breadcrumbs a, [class*=breadcrumb] a, nav[aria-label*=readcrumb] a")
        cats = [clean_text(a.get_text()) for a in crumbs[1:]]
        cats = [c for c in cats if not GENERIC_CRUMB_RE.match(c)]
    if not cats and ld and ld.get("category"):
        cats = [clean_text(c) for c in _as_list(ld["category"]) if isinstance(c, str)]
    rec["categories"] = " | ".join(_dedupe([c for c in cats if c and c != rec["name"]]))

    # ---- descriptions
    sd = first(soup, [".woocommerce-product-details__short-description", ".short-description",
                      "div[itemprop=description]"])
    rec["short_description"] = clean_text(sd) if sd else ""
    desc_el = first(soup, ["#tab-description", ".woocommerce-Tabs-panel--description",
                           ".product-description", "#description", ".product-details-content",
                           ".product-content", "div[itemprop=description]", "[id*=description]",
                           "[class*=product-desc]", "#review", "#introduction"])
    rec["description"] = (clean_text(desc_el) if desc_el else "") \
        or clean_text((ld or {}).get("description")) \
        or rec["short_description"] or clean_text(meta.get("description", ""))
    rec["description_html"] = sanitize_html(desc_el, url)

    # ---- tags / meta
    tags = [clean_text(a.get_text()) for a in soup.select(".product_meta .tagged_as a, .tagged_as a")]
    rec["tags"] = " | ".join(_dedupe(t for t in tags if t))
    title = soup.find("title")
    rec["meta_title"] = clean_text(title.get_text()) if title else clean_text(meta.get("og:title", ""))
    rec["meta_description"] = clean_text(meta.get("description") or meta.get("og:description", ""))

    # ---- specs: WooCommerce attributes, JSON-LD, then tables/lists in tabs & summary
    specs = {}
    for tr in soup.select("table.woocommerce-product-attributes tr, table.shop_attributes tr"):
        th, td = tr.find("th"), tr.find("td")
        if th and td:
            specs[clean_text(th.get_text())] = clean_text(td.get_text(" "))
    for prop in _as_list((ld or {}).get("additionalProperty")):
        if isinstance(prop, dict) and prop.get("name"):
            specs.setdefault(clean_text(prop["name"]), clean_text(prop.get("value")))
    panels = soup.select(".woocommerce-Tabs-panel, .wc-tab, .product-specs, .product-attributes, "
                         ".specifications, #specifications, .product-features, #specs, [id*=spec], "
                         "[class*=spec], [class*=attribute], [class*=feature]")
    panels = [p for p in panels if p.name not in ("html", "body")]
    for k, v in extract_specs(desc_el, *panels, sd).items():
        specs.setdefault(k, v)
    rec["specs"] = json.dumps(specs, ensure_ascii=False) if specs else ""

    # ---- images: gallery first, then JSON-LD, og:image; images inside the description kept separately
    imgs = []
    for a in soup.select(".woocommerce-product-gallery__image a[href], .woocommerce-product-gallery a[href]"):
        if re.search(r"\.(jpe?g|png|webp|gif|avif)(?:$|\?)", a["href"], re.I):
            imgs.append(a["href"])
    for img in soup.select(".woocommerce-product-gallery img, .product-gallery img, "
                           ".product-images img, [class*=product-image] img, [id*=allery] img, "
                           "[class*=gallery] img, main .swiper img, section .swiper img"):
        u = best_img_src(img)
        if u:
            imgs.append(u)
    imgs += _image_urls((ld or {}).get("image"))
    if meta.get("og:image"):
        imgs.append(meta["og:image"])
    set_images(rec, [urljoin(url, u) for u in imgs if isinstance(u, str)][: MAX_IMAGES_PER_PRODUCT * 2])
    gallery = set(rec["images"].split(" | "))
    rec["description_images"] = " | ".join(u for u in content_images(desc_el, url) if u not in gallery)
    rec["files"] = " | ".join(find_files(soup, url))
    return rec


# --------------------------------------------------------------------------- #
# Fast product-page parser (lxml + XPath; ~10x less CPU than BeautifulSoup/soupsieve)
# --------------------------------------------------------------------------- #

import copy
import functools
from lxml import etree

_CSS_TOKEN = re.compile(r"""
    (?P<tag>^[a-zA-Z][\w-]*|\*)?
""", re.X)


def _css_compound_to_xpath(comp: str) -> str:
    """'div.a#b[x*=y]:not([c*=d]):last-child' -> 'div[...][...]' (subset used in this file)."""
    m = re.match(r"^([a-zA-Z][\w-]*|\*)?", comp)
    tag = m.group(1) or "*"
    rest = comp[m.end():]
    preds = []
    while rest:
        if rest.startswith("."):
            mm = re.match(r"^\.([\w-]+)", rest)
            preds.append(f"contains(concat(' ',normalize-space(@class),' '),' {mm.group(1)} ')")
        elif rest.startswith("#"):
            mm = re.match(r"^#([\w-]+)", rest)
            preds.append(f"@id='{mm.group(1)}'")
        elif rest.startswith("["):
            mm = re.match(r"^\[([\w-]+)(?:([*^$]?=)['\"]?([^'\"\]]*)['\"]?)?\]", rest)
            attr, op, val = mm.group(1), mm.group(2), mm.group(3)
            if not op:
                preds.append(f"@{attr}")
            elif op == "=":
                preds.append(f"@{attr}='{val}'")
            elif op == "*=":
                preds.append(f"contains(@{attr},'{val}')")
            elif op == "^=":
                preds.append(f"starts-with(@{attr},'{val}')")
            else:
                preds.append(f"substring(@{attr},string-length(@{attr})-{len(val) - 1})='{val}'")
        elif rest.startswith(":not("):
            depth, i = 0, 0
            for i, ch in enumerate(rest):
                depth += ch == "("
                depth -= ch == ")"
                if depth == 0 and ch == ")":
                    break
            inner = rest[5:i]
            inner_x = _css_compound_to_xpath(inner)
            inner_preds = inner_x[inner_x.index("["):] if "[" in inner_x else ""
            preds.append(f"not(self::*{inner_preds})")
            mm = re.match(r"^.{%d}" % (i + 1), rest, re.S)
        elif rest.startswith(":last-child"):
            mm = re.match(r"^:last-child", rest)
            preds.append("not(following-sibling::*)")
        elif rest.startswith(":first-child"):
            mm = re.match(r"^:first-child", rest)
            preds.append("not(preceding-sibling::*)")
        else:
            raise ValueError(f"unsupported selector: {comp}")
        rest = rest[mm.end():]
    return tag + "".join(f"[{p}]" for p in preds)


@functools.lru_cache(maxsize=512)
def css_xpath(selector: str) -> etree.XPath:
    parts = []
    for group in re.split(r",(?![^\[(]*[\])])", selector):
        tokens = re.findall(r">|[^\s>]+(?:\([^)]*\))?[^\s>]*", group.strip())
        xp, axis = "", "descendant-or-self::"
        for t in tokens:
            if t == ">":
                axis = "/"
                continue
            step = _css_compound_to_xpath(t)
            if not xp:
                xp = f"descendant-or-self::{step}"
            else:
                xp += f"/{step}" if axis == "/" else f"//{step}"
            axis = "//"
        parts.append(xp)
    return etree.XPath(" | ".join(parts))


def lx_select(root, selector: str) -> list:
    return css_xpath(selector)(root)


def lx_one(root, selector: str):
    r = css_xpath(selector)(root)
    return r[0] if r else None


def lx_first(root, selectors: list[str]):
    for sel in selectors:
        el = lx_one(root, sel)
        if el is not None:
            return el
    return None


def _ws(s: str) -> str:
    s = (s or "").replace("\xa0", " ")
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n", s)
    return "\n".join(line.strip() for line in s.split("\n")).strip()


def lx_inline(el) -> str:
    """Plain text of an element on one line."""
    if el is None:
        return ""
    return re.sub(r"\s+", " ", html_lib.unescape(" ".join(el.itertext()))).strip()


_SKIP_TEXT_TAGS = {"script", "style", "noscript", "template"}


def lx_text(el) -> str:
    """Readable multi-line text: newlines only around block elements, ' | ' between table cells."""
    if el is None:
        return ""
    parts: list[str] = []

    def walk(e, root=False):
        tag = e.tag if isinstance(e.tag, str) else None
        if tag and tag not in _SKIP_TEXT_TAGS:
            if tag in BLOCK_TAGS:
                parts.append("\n")
            if tag == "br":
                parts.append("\n")
            if e.text:
                parts.append(e.text)
            for c in e:
                walk(c)
            if tag in ("td", "th"):
                parts.append(" | ")
            if tag in BLOCK_TAGS:
                parts.append("\n")
        if not root and e.tail:
            parts.append(e.tail)

    walk(el, root=True)
    text = re.sub(r"[ \t]*\|[ \t]*\n", "\n", "".join(parts))
    return _ws(html_lib.unescape(text))


def lx_sanitize(el, base_url: str) -> str:
    if el is None:
        return ""
    el = copy.deepcopy(el)
    for bad in el.xpath(".//script|.//style|.//noscript|.//iframe|.//form|.//button|.//svg|.//input|"
                        ".//select|.//textarea|.//link|.//meta|.//template|.//comment()"):
        if bad.getparent() is not None:
            bad.drop_tree() if isinstance(bad.tag, str) else bad.getparent().remove(bad)
    for img in list(el.iter("img")):
        src = best_img_src(img)
        if src:
            img.set("src", urljoin(base_url, src))
        elif img.getparent() is not None:
            img.drop_tree()
    for tag in list(el.iterdescendants()):
        if not isinstance(tag.tag, str):
            continue
        if tag.tag not in SAFE_TAGS:
            tag.drop_tag()
            continue
        for k in list(tag.attrib):
            if k not in SAFE_ATTRS:
                del tag.attrib[k]
        for k in ("href", "src"):
            if tag.get(k):
                tag.set(k, urljoin(base_url, tag.get(k)))
    html = (el.text or "") + "".join(lxml.html.tostring(c, encoding="unicode") for c in el)
    html = re.sub(r"\n\s*\n+", "\n", html)
    return re.sub(r"<(p|div|span)>\s*</\1>", "", html).strip()


def lx_content_images(el, base_url: str) -> list[str]:
    if el is None:
        return []
    out = []
    for img in el.iter("img"):
        u = best_img_src(img)
        if u and not u.lower().endswith(".svg"):
            out.append(urljoin(base_url, u))
    return _dedupe(out)


def lx_specs(*containers) -> dict:
    """Key/value pairs from 2-column tables, <dl> lists and list items ('key: value' or two child elements)."""
    specs: dict[str, str] = {}

    def add(k, v):
        k, v = _ws(k).strip(" :：-"), _ws(v)
        if k and v and len(k) <= 60 and k != v and k not in specs:
            specs[k] = v[:1000]

    seen = set()
    for c in containers:
        if c is None or id(c) in seen:
            continue
        seen.add(id(c))
        for tr in c.iter("tr"):
            cells = [x for x in tr if x.tag in ("th", "td")]
            if len(cells) == 2:
                add(lx_inline(cells[0]), lx_inline(cells[1]))
        for dt in c.iter("dt"):
            dd = dt.getnext()
            while dd is not None and dd.tag != "dd":
                dd = dd.getnext() if dd.tag not in ("dt",) else None
            if dd is not None:
                add(lx_inline(dt), lx_inline(dd))
        for lst in c.iter("ul", "ol"):
            pairs = []
            for li in lst:
                if li.tag != "li":
                    continue
                kids = [k for k in li if isinstance(k.tag, str) and lx_inline(k)]
                if len(kids) == 2 and kids[0].find(".//ul") is None and kids[0].find(".//table") is None:
                    pairs.append((lx_inline(kids[0]), lx_inline(kids[1])))
                    continue
                t = lx_inline(li)
                m = re.match(r"^([^:：]{2,40})\s*[:：]\s*(.{1,300})$", t, re.S)
                if m:
                    pairs.append(m.groups())
            if len(pairs) >= 2:
                for k, v in pairs:
                    add(k, v)
    return specs


_BLOCKY = {"div", "li", "p", "table", "ul", "ol", "section", "article"}


def lx_labeled_values(root) -> dict:
    out = {}
    for el in root.iter("div", "li", "p", "span", "td"):
        if any(isinstance(d.tag, str) and d.tag in _BLOCKY for d in el.iterdescendants()):
            continue
        t = lx_inline(el)
        if len(t) > 160:
            continue
        m = re.match(r"^([^:：]{2,30})\s*[:：]\s*(.{1,120})$", t)
        if m and m.group(1).strip() not in out:
            out[m.group(1).strip()] = m.group(2).strip()
    return out


def _in_product_card(el) -> bool:
    """True for elements inside related/upsell/listing product cards (not the main product)."""
    for a in el.iterancestors():
        cls = f" {a.get('class', '')} "
        if re.search(r"\s(related|upsells?|cross-sells?|product-item|products-list|product-card|product-box|"
                     r"related-products|products|product-detail-area|swiper[\w-]*|slick[\w-]*|owl-carousel|"
                     r"carousel|widget|sidebar|product-grid|product-list)\s", cls) \
                or (a.tag == "li" and " product " in cls) or (a.tag == "article" and " product " in cls):
            return True
    return False


def _lx_jsonld(root):
    for sc in root.xpath("//script[contains(translate(@type,'LDJSON','ldjson'),'ld+json')]"):
        raw = re.sub(r"[\x00-\x1f]", " ", sc.text or "").strip()
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except ValueError:
            continue
        stack = [data]
        while stack:
            node = stack.pop()
            if isinstance(node, list):
                stack.extend(node)
            elif isinstance(node, dict):
                yield node
                for k in ("@graph", "mainEntity", "itemListElement", "item"):
                    if k in node:
                        stack.append(node[k])


def lx_parse(html: str):
    """Parse and strip menus/scripts/icons in place. Returns the lxml root or None."""
    try:
        root = lxml.html.fromstring(html.encode("utf-8", "replace"), parser=_LXML_PARSER)
    except Exception:
        return None
    if root is None:
        return None
    drop = []
    for el in root.iter("script", "style", "noscript", "svg", "iframe", "template", "link"):
        if el.tag == "script" and "json" in (el.get("type") or "").lower():
            continue
        drop.append(el)
    for el in root.iter("header", "footer", "nav", "aside", "div", "ul", "section"):
        if el.tag in ("div", "ul", "section"):
            if not _JUNK_CLASS_RE.search(f"{el.get('class', '')} {el.get('id', '')}"):
                continue
        elif el.tag == "nav" and re.search(r"breadcrumb", f"{el.get('aria-label', '')} {el.get('class', '')}", re.I):
            continue
        if any(a.tag in ("main", "article") for a in el.iterancestors()):
            continue
        if el.find(".//h1") is not None or el.xpath(".//*[contains(@class,'product_title') or contains(@class,'breadcrumb')]"):
            continue
        drop.append(el)
    for el in drop:
        if el.getparent() is not None:
            el.drop_tree()
    return root


def parse_product_page(html: str, url: str, site: str, trust_as_product: bool = False, root=None) -> dict | None:
    root = root if root is not None else lx_parse(html)
    if root is None:
        return None
    ld, crumbs_ld = None, []
    for node in _lx_jsonld(root):
        if _is_type(node, "Product") or _is_type(node, "ProductGroup"):
            if ld is None or (node.get("offers") and not ld.get("offers")):
                ld = node
        elif _is_type(node, "BreadcrumbList") and not crumbs_ld:
            items = sorted((x for x in _as_list(node.get("itemListElement")) if isinstance(x, dict)),
                           key=lambda x: x.get("position") or 0)
            crumbs_ld = [clean_text(x.get("name") or (x.get("item") or {}).get("name") if isinstance(x.get("item"), dict)
                                    else x.get("name")) for x in items]
    meta = {}
    for m in root.iter("meta"):
        k = (m.get("property") or m.get("name") or m.get("itemprop") or "").strip().lower()
        if k and m.get("content") and k not in meta:
            meta[k] = m.get("content").strip()
    body = root.find(".//body")
    body_cls = body.get("class", "") if body is not None else ""

    html_signals = lx_one(root, "form.cart, .product_title, .single-product-summary, "
                                "[itemtype*='schema.org/Product'], .woocommerce-product-gallery") is not None
    is_product = (ld is not None or meta.get("og:type", "").lower() == "product"
                  or "single-product" in body_cls or html_signals)
    generic = not is_product and (trust_as_product or looks_like_product_url(url))
    if not is_product and not generic:
        return None

    rec = {k: "" for k in FIELDS}
    rec.update(site=site, source="page", url=url,
               crawled_at=datetime.now(timezone.utc).isoformat(timespec="seconds"))

    # ---- name
    h1 = lx_first(root, ["h1.product_title", "h1.product-title", ".product_title", "main h1", "h1"])
    crumbs_last = lx_one(root, "nav[aria-label*=readcrumb] li:last-child, [class*=breadcrumb] li:last-child")
    title_el = root.find(".//title")
    title_name = ""
    if title_el is not None:
        bits = re.split(r"\s+[-|–—]\s+", lx_inline(title_el))
        title_name = bits[0] if bits else ""
    rec["name"] = clean_text((ld or {}).get("name")) or lx_inline(h1) or clean_text(meta.get("og:title", "")) \
        or lx_inline(crumbs_last) or title_name
    if not rec["name"]:
        return None
    labeled = lx_labeled_values(root) if generic else {}

    # ---- id / sku
    rec["sku"] = clean_text((ld or {}).get("sku") or (ld or {}).get("mpn") or "")
    if not rec["sku"]:
        el = lx_one(root, ".sku, [itemprop=sku]")
        if el is not None:
            rec["sku"] = clean_text(el.get("content") or lx_inline(el))
        else:
            rec["sku"] = next((v for k, v in labeled.items() if SKU_LABEL_RE.search(k)), "")
    m = re.search(r"postid-(\d+)", body_cls)
    if m:
        rec["product_id"] = m.group(1)
    else:
        el = lx_one(root, "[data-product_id], input[name=add-to-cart], button[name=add-to-cart]")
        if el is not None:
            rec["product_id"] = el.get("data-product_id") or el.get("value") or ""
    main_box = None
    if rec["product_id"]:
        main_box = lx_one(root, f"#product-{rec['product_id']}")
    if main_box is None:
        main_box = next((d for d in lx_select(root, "div[id^=product-]") if " product " in f" {d.get('class', '')} "), None)

    # ---- brand
    brand = (ld or {}).get("brand")
    if isinstance(brand, list):
        brand = brand[0] if brand else ""
    if isinstance(brand, dict):
        brand = brand.get("name", "")
    rec["brand"] = clean_text(brand) or next((v for k, v in labeled.items() if BRAND_LABEL_RE.search(k)), "")

    # ---- price (JSON-LD offers first, then the main product's price box, never a related product's)
    prices, currency, availability = [], "", ""
    for off in _as_list((ld or {}).get("offers")):
        if not isinstance(off, dict):
            continue
        currency = currency or off.get("priceCurrency", "")
        availability = availability or str(off.get("availability", ""))
        for key in ("price", "lowPrice", "highPrice"):
            v = off.get(key)
            if v not in (None, ""):
                n = parse_number(str(v))
                if n is not None:
                    prices.append(n)
        for ps in _as_list(off.get("priceSpecification")):
            if isinstance(ps, dict) and ps.get("price") not in (None, ""):
                n = parse_number(str(ps["price"]))
                if n is not None:
                    prices.append(n)
                currency = currency or ps.get("priceCurrency", "")

    price_el = None
    for sel in (".summary .price", ".entry-summary .price", ".product-info .price", ".price-wrapper .price",
                "p.price", "[itemprop=price]", ".product-price", "span.price"):
        price_el = next((e for e in lx_select(root, sel) if not _in_product_card(e)), None)
        if price_el is not None:
            break
    if price_el is None and generic:
        price_el = next((e for e in lx_select(root, "[class*=price], [id*=price]")
                         if not _in_product_card(e) and extract_numbers(lx_inline(e)) and len(lx_inline(e)) < 120), None)
    price_text = ""
    if price_el is not None:
        price_text = clean_text(price_el.get("content") or lx_inline(price_el))
        ins, dels = price_el.find(".//ins"), price_el.find(".//del")
        if ins is not None and dels is not None:
            sale, reg = extract_numbers(lx_inline(ins)), extract_numbers(lx_inline(dels))
            if sale:
                rec["sale_price"] = sale[0]
            if reg:
                rec["regular_price"] = reg[0]
    if not price_text and meta.get("product:price:amount"):
        price_text = meta["product:price:amount"]
    rec["price_text"] = price_text
    if not prices and price_text:
        nums = [n for n in extract_numbers(price_text) if n > 0]
        prices = [rec["sale_price"]] if rec["sale_price"] != "" else nums
    if prices:
        rec["min_price"], rec["max_price"] = min(prices), max(prices)
        rec["price"] = rec["sale_price"] if rec["sale_price"] != "" else min(prices)
    rec["currency"] = currency or meta.get("product:price:currency", "") or detect_currency(price_text)

    # ---- stock
    av = availability.lower()
    stock_el = next((e for e in lx_select(root, ".summary .stock, p.stock, .stock") if not _in_product_card(e)), None)
    stock_txt = lx_inline(stock_el)
    stock_cls = stock_el.get("class", "") if stock_el is not None else ""
    box_cls = f" {main_box.get('class', '')} " if main_box is not None else ""
    if "outofstock" in av or "out-of-stock" in stock_cls or "ناموجود" in stock_txt or " outofstock " in box_cls:
        rec["in_stock"] = False
    elif "instock" in av or "in-stock" in stock_cls or "موجود" in stock_txt or " instock " in box_cls:
        rec["in_stock"] = True

    # ---- descriptions (needed before specs)
    sd = lx_first(root, [".woocommerce-product-details__short-description", ".short-description",
                         "div[itemprop=description]"])
    rec["short_description"] = lx_text(sd)
    desc_el = lx_first(root, ["#tab-description", ".woocommerce-Tabs-panel--description",
                              ".product-description", "#description", ".product-details-content",
                              ".product-content", "div[itemprop=description]", "[id*=description]",
                              "[class*=product-desc]", "#review", "#introduction"])
    rec["description"] = lx_text(desc_el) or clean_text((ld or {}).get("description")) \
        or rec["short_description"] or clean_text(meta.get("description", ""))
    rec["description_html"] = lx_sanitize(desc_el, url)

    # ---- specs
    specs = {}
    for tr in lx_select(root, "table.woocommerce-product-attributes tr, table.shop_attributes tr"):
        th, td = tr.find("th"), tr.find("td")
        if th is not None and td is not None:
            specs[lx_inline(th)] = lx_inline(td)
    for prop in _as_list((ld or {}).get("additionalProperty")):
        if isinstance(prop, dict) and prop.get("name"):
            specs.setdefault(clean_text(prop["name"]), clean_text(prop.get("value")))
    panels = [p for p in lx_select(root, ".woocommerce-Tabs-panel, .wc-tab, .product-specs, .product-attributes, "
                                         ".specifications, #specifications, .product-features, #specs, [id*=spec], "
                                         "[class*=spec], [class*=attribute], [class*=feature]")
              if p.tag not in ("html", "body") and not _in_product_card(p)]
    for k, v in lx_specs(desc_el, *panels, sd).items():
        specs.setdefault(k, v)
    rec["specs"] = json.dumps(specs, ensure_ascii=False) if specs else ""
    if not rec["brand"]:
        rec["brand"] = next((v for k, v in specs.items() if BRAND_LABEL_RE.search(k)), "")

    # ---- categories
    cats = [lx_inline(a) for a in lx_select(root, ".product_meta .posted_in a, .posted_in a")]
    if not cats:
        crumbs = lx_select(root, ".woocommerce-breadcrumb a, nav.rank-math-breadcrumb a, .breadcrumb a, "
                                 ".breadcrumbs a, [class*=breadcrumb] a, nav[aria-label*=readcrumb] a")
        cats = [lx_inline(a) for a in crumbs[1:]]
    if not cats and crumbs_ld:
        cats = crumbs_ld[1:]
    if not cats and ld and ld.get("category"):
        cats = [clean_text(c) for c in _as_list(ld["category"]) if isinstance(c, str)]
    if not cats:
        cats = [v for k, v in specs.items() if re.match(r"^(دسته ?بندی|دسته‌بندی|category)$", k, re.I)]
    rec["categories"] = " | ".join(_dedupe([c for c in cats if c and c != rec["name"] and not GENERIC_CRUMB_RE.match(c)]))

    # ---- tags / meta
    rec["tags"] = " | ".join(_dedupe(lx_inline(a) for a in lx_select(root, ".product_meta .tagged_as a, .tagged_as a")))
    rec["meta_title"] = lx_inline(title_el) if title_el is not None else clean_text(meta.get("og:title", ""))
    rec["meta_description"] = clean_text(meta.get("description") or meta.get("og:description", ""))

    # ---- images
    imgs = []
    for a in lx_select(root, ".woocommerce-product-gallery__image a[href], .woocommerce-product-gallery a[href]"):
        if re.search(r"\.(jpe?g|png|webp|gif|avif)(?:$|\?)", a.get("href", ""), re.I):
            imgs.append(a.get("href"))
    for img in lx_select(root, ".woocommerce-product-gallery img, .product-gallery img, .product-images img, "
                               "[class*=product-image] img, [id*=allery] img, [class*=gallery] img, "
                               "main .swiper img, section .swiper img"):
        if _in_product_card(img):
            continue
        u = best_img_src(img)
        if u:
            imgs.append(u)
    imgs += _image_urls((ld or {}).get("image"))
    if meta.get("og:image"):
        imgs.append(meta["og:image"])
    set_images(rec, [urljoin(url, u) for u in imgs if isinstance(u, str)][: MAX_IMAGES_PER_PRODUCT * 2])
    gallery = set(rec["images"].split(" | "))
    rec["description_images"] = " | ".join(u for u in lx_content_images(desc_el, url) if u not in gallery)

    # ---- files
    scope = lx_first(root, ["main", "#main", ".site-main", ".product", "article"])
    scope = scope if scope is not None else root
    rec["files"] = " | ".join(_dedupe(urljoin(url, a.get("href")) for a in scope.iter("a")
                                      if a.get("href") and FILE_EXT_RE.search(a.get("href"))))
    return rec


# --------------------------------------------------------------------------- #
# Price lists: products as rows of an HTML table or items of JSON embedded in the page
# --------------------------------------------------------------------------- #

NAME_KEYS = ["name", "title", "product_name", "productName", "product", "model", "code", "sku",
             "نام", "نام کالا", "کد کالا", "مدل", "عنوان"]
PRICE_KEYS = ["price", "final_price", "finalPrice", "sale_price", "salePrice", "regular_price", "unit_price",
              "قیمت", "amount", "cost"]
NAME_HDR_RE = re.compile(r"نام|مدل|کد|عنوان|شرح|کالا|محصول|name|model|code|product|item|description", re.I)
PRICE_HDR_RE = re.compile(r"قیمت|بها|price|ریال|تومان|cost", re.I)
SKIP_COL_RE = re.compile(r"^(ردیف|#|row|no\.?|شماره)$", re.I)


def _listing_record(site, page_url, name, price_raw, specs, category="", link="", image="", currency="", source="table"):
    rec = {k: "" for k in FIELDS}
    rec.update(site=site, source=source, name=clean_text(name)[:300],
               crawled_at=datetime.now(timezone.utc).isoformat(timespec="seconds"))
    anchor = re.sub(r"\s+", "-", f"{category}-{rec['name']}".strip("-"))[:120]
    rec["url"] = urljoin(page_url, link) if link else f"{page_url.split('#')[0]}#{anchor}"
    if isinstance(price_raw, (int, float)) and not isinstance(price_raw, bool):
        rec["price"] = rec["min_price"] = rec["max_price"] = price_raw
        rec["price_text"] = str(price_raw)
    else:
        txt = clean_text(price_raw)
        rec["price_text"] = txt
        nums = [n for n in extract_numbers(txt) if n > 0]
        if nums:
            rec["price"], rec["min_price"], rec["max_price"] = nums[0], min(nums), max(nums)
        currency = currency or detect_currency(txt)
    rec["currency"] = currency
    rec["categories"] = clean_text(category)
    specs = {clean_text(k): clean_text(v) for k, v in specs.items() if clean_text(k) and clean_text(v)}
    rec["specs"] = json.dumps(specs, ensure_ascii=False) if specs else ""
    if image:
        set_images(rec, [urljoin(page_url, image)])
    return rec


def records_from_json(data, page_url: str, site: str) -> list[dict]:
    """Find lists of product-like objects (a name key and a price key) anywhere in a JSON document."""
    out = []

    def pick(keys, candidates):
        lower = {k.lower(): k for k in keys}
        for c in candidates:
            if c.lower() in lower:
                return lower[c.lower()]
        return None

    def scalar(v):
        if isinstance(v, dict):
            v = v.get("value", v.get("label", ""))
        if isinstance(v, (list, dict)) or v is None:
            return ""
        return v

    def walk(node, parent, depth):
        if depth > 12:
            return
        if isinstance(node, list):
            dicts = [x for x in node if isinstance(x, dict)]
            if len(dicts) >= 2:
                keys = set().union(*(d.keys() for d in dicts))
                nk, pk = pick(keys, NAME_KEYS), pick(keys, PRICE_KEYS)
                if nk and pk and sum(1 for d in dicts if d.get(nk) not in (None, "")) >= 0.6 * len(dicts):
                    labels = {}
                    if isinstance(parent, dict):
                        for c in parent.get("columns") or []:
                            if isinstance(c, dict) and c.get("key"):
                                labels[c["key"]] = c.get("label") or c["key"]
                    category = ""
                    if isinstance(parent, dict):
                        category = parent.get("title") or parent.get("name") or parent.get("category") or ""
                        category = category if isinstance(category, str) else ""
                    for d in dicts:
                        if d.get(nk) in (None, ""):
                            continue
                        specs = {}
                        for k, v in d.items():
                            if k in (nk, pk) or str(k).lower() in ("row", "id", "url", "link", "permalink",
                                                                    "slug", "image", "img", "thumbnail"):
                                continue
                            val = scalar(v)
                            if val != "":
                                label = labels.get(k) or (v.get("label") if isinstance(v, dict) else None) or k
                                specs[str(label)] = str(val)
                        link = d.get("url") or d.get("permalink") or d.get("link") or ""
                        img = d.get("image") or d.get("img") or d.get("thumbnail") or ""
                        out.append(_listing_record(site, page_url, str(d[nk]), d.get(pk), specs, category,
                                                   link if isinstance(link, str) else "",
                                                   img if isinstance(img, str) else "", source="json"))
                    return
            for x in node:
                walk(x, parent, depth + 1)
        elif isinstance(node, dict):
            for v in node.values():
                if isinstance(v, (list, dict)):
                    walk(v, node, depth + 1)

    walk(data, None, 0)
    return out


def _table_category(table) -> str:
    cap = table.find("caption")
    if cap and clean_text(cap.get_text()):
        return clean_text(cap.get_text())
    prev = table.find_previous(["h2", "h3", "h4", "h5"])
    return clean_text(prev.get_text()) if prev else ""


def records_from_table(table, page_url: str, site: str, require_price: bool = True) -> list[dict]:
    """Each data row of a price table becomes a product. Handles repeated header rows and
    single-cell group rows (used as the category)."""
    out = []
    header: list[str] = []
    group = ""
    base_cat = _table_category(table)
    for tr in table.find_all("tr"):
        cells = tr.find_all(["th", "td"], recursive=False)
        if not cells:
            continue
        texts = [clean_text(c.get_text(" ")) for c in cells]
        nonempty = [t for t in texts if t]
        if len(nonempty) <= 1 and (len(cells) == 1 or (cells[0].get("colspan") or "1") != "1"):
            group = nonempty[0] if nonempty else group            # group title row
            continue
        keyword_cells = sum(1 for t in texts if NAME_HDR_RE.search(t) or PRICE_HDR_RE.search(t))
        numeric_cells = sum(1 for t in texts if re.fullmatch(r"[\d\s,٬.۰-۹/×x\-]+", t or "x"))
        is_header = all(c.name == "th" for c in cells) or texts == header or (
            keyword_cells >= 2 and numeric_cells == 0)
        if is_header:
            header = texts
            continue
        if not header or len(texts) < 2:
            continue
        cols = header + [f"ستون {i + 1}" for i in range(len(header), len(texts))]
        price_i = next((i for i, h in enumerate(cols) if PRICE_HDR_RE.search(h)), None)
        if price_i is None and require_price:
            continue
        name_i = next((i for i, h in enumerate(cols) if NAME_HDR_RE.search(h) and i != price_i
                       and not SKIP_COL_RE.match(h)), None)
        if name_i is None:
            name_i = next((i for i, t in enumerate(texts) if t and not extract_numbers(t) and i != price_i), 0)
        name = texts[name_i] if name_i < len(texts) else ""
        if not name:
            continue
        specs = {cols[i]: t for i, t in enumerate(texts)
                 if i not in (name_i, price_i) and t and not SKIP_COL_RE.match(cols[i])}
        link = ""
        a = tr.find("a", href=True)
        if a and not a["href"].startswith(("#", "tel:", "mailto:", "javascript:")):
            link = a["href"]
        img = tr.find("img")
        price_txt = texts[price_i] if price_i is not None and price_i < len(texts) else ""
        currency = detect_currency(" ".join(cols)) if price_i is not None else ""
        category = " / ".join(x for x in (base_cat, group) if x)
        out.append(_listing_record(site, page_url, name, price_txt, specs, category, link,
                                   best_img_src(img) if img else "", currency, source="table"))
    return out


def extract_listing_records(html: str, page_url: str, site: str, require_price: bool = True) -> list[dict]:
    """Products from price tables and from product lists embedded as JSON in the page."""
    soup = BeautifulSoup(preclean(html), "lxml")
    out = []
    for script in soup.find_all("script"):
        typ = (script.get("type") or "").lower()
        if "json" not in typ or "ld+json" in typ:
            continue
        text = script.string or script.get_text() or ""
        if len(text) < 100:
            continue
        try:
            out += records_from_json(json.loads(text), page_url, site)
        except ValueError:
            continue
    for table in soup.find_all("table"):
        if len(table.find_all("tr")) >= 3:
            out += records_from_table(table, page_url, site, require_price)
    return out


def merge_records(base: dict, extra: dict | None) -> dict:
    """Fill gaps in an API record with what the product page adds (images, specs, files, full content…)."""
    if not extra:
        return base
    out = dict(base)
    for k in FIELDS:
        if k in ("source", "url", "crawled_at", "specs", "images", "image_main", "image_count"):
            continue
        if out.get(k) in ("", None) and extra.get(k) not in ("", None):
            out[k] = extra[k]
    if len(extra.get("description_html", "")) > len(out.get("description_html", "")) * 1.2:
        out["description_html"] = extra["description_html"]
        out["description"] = extra.get("description") or out.get("description", "")
    specs = json.loads(out["specs"]) if out.get("specs") else {}
    for k, v in (json.loads(extra["specs"]) if extra.get("specs") else {}).items():
        specs.setdefault(k, v)
    out["specs"] = json.dumps(specs, ensure_ascii=False) if specs else ""
    imgs = [u for u in (out.get("images") or "").split(" | ") if u]
    imgs += [u for u in (extra.get("images") or "").split(" | ") if u]
    set_images(out, imgs)
    out["source"] = f"{base.get('source', '')}+page"
    return out


def product_from_store_api(item: dict, site: str) -> dict:
    rec = {k: "" for k in FIELDS}
    prices = item.get("prices") or {}
    try:
        mu = int(prices.get("currency_minor_unit") or 0)
    except (TypeError, ValueError):
        mu = 0

    def money(v):
        if v in (None, ""):
            return ""
        try:
            return int(v) / (10 ** mu) if mu else int(v)
        except (TypeError, ValueError):
            n = parse_number(str(v))
            return n if n is not None else ""

    price, reg, sale = money(prices.get("price")), money(prices.get("regular_price")), money(prices.get("sale_price"))
    rng = prices.get("price_range") or {}
    mn, mx = money(rng.get("min_amount")), money(rng.get("max_amount"))
    if sale != "" and sale == reg:
        sale = ""

    url = item.get("permalink", "") or ""
    attrs = {}
    for a in item.get("attributes") or []:
        terms = [t.get("name", "") for t in (a.get("terms") or []) if isinstance(t, dict)]
        attrs[clean_text(a.get("name"))] = "، ".join(clean_text(t) for t in terms)
    desc_raw = item.get("description") or ""
    try:
        desc_el = lxml.html.fragment_fromstring(desc_raw, create_parent="div") if desc_raw.strip() else None
    except Exception:
        desc_el = None
    for k, v in lx_specs(desc_el).items():
        attrs.setdefault(k, v)

    brand = ""
    for b in item.get("brands") or []:
        if isinstance(b, dict) and b.get("name"):
            brand = clean_text(b["name"])
            break

    rec.update(
        site=site, source="store_api",
        url=item.get("permalink", ""),
        product_id=item.get("id", ""),
        name=clean_text(item.get("name")),
        sku=clean_text(item.get("sku")),
        price=price, regular_price=reg, sale_price=sale,
        min_price=mn if mn != "" else price, max_price=mx if mx != "" else price,
        currency=prices.get("currency_code") or prices.get("currency_symbol") or "",
        price_text=clean_text(item.get("price_html")),
        in_stock=item.get("is_in_stock", ""),
        categories=" | ".join(clean_text(c.get("name")) for c in item.get("categories") or [] if isinstance(c, dict)),
        brand=brand,
        short_description=clean_text(item.get("short_description")),
        description=lx_text(desc_el) if desc_el is not None else "",
        description_html=lx_sanitize(desc_el, url) if desc_el is not None else "",
        specs=json.dumps(attrs, ensure_ascii=False) if attrs else "",
        tags=" | ".join(clean_text(t.get("name")) for t in item.get("tags") or [] if isinstance(t, dict)),
        description_images=" | ".join(lx_content_images(desc_el, url)),
        crawled_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
    )
    set_images(rec, [i.get("src", "") for i in item.get("images") or [] if isinstance(i, dict)])
    return rec


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #

class JsonlWriter:
    def __init__(self, path: str, resume: bool):
        self.path = path
        self.lock = threading.Lock()
        self.seen_urls: set[str] = set()
        self.count = 0
        if resume and os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                for line in f:
                    try:
                        self.seen_urls.add(json.loads(line)["url"])
                        self.count += 1
                    except (ValueError, KeyError):
                        pass
            self.fh = open(path, "a", encoding="utf-8")
        else:
            self.fh = open(path, "w", encoding="utf-8")

    def write(self, rec: dict) -> bool:
        key = rec.get("url") or f"id:{rec.get('product_id')}"
        with self.lock:
            if key in self.seen_urls:
                return False
            self.seen_urls.add(key)
            self.fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
            self.fh.flush()
            self.count += 1
            return True

    def close(self):
        self.fh.close()


def iter_jsonl(src):
    """Yield records from a .jsonl path (streamed, one line at a time) or from a list."""
    if isinstance(src, (list, tuple)):
        yield from src
        return
    if not src or not os.path.exists(src):
        return
    with open(src, encoding="utf-8") as f:
        for line in f:
            try:
                yield json.loads(line)
            except ValueError:
                pass


def read_jsonl(path: str) -> list[dict]:
    return list(iter_jsonl(path))


def write_csv(src, path: str) -> int:
    n = 0
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        for r in iter_jsonl(src):
            w.writerow({k: r.get(k, "") for k in FIELDS})
            n += 1
    return n


def write_json(src, path: str) -> int:
    """JSON array written record by record (never holds the whole list in memory)."""
    n = 0
    with open(path, "w", encoding="utf-8") as f:
        f.write("[\n")
        for r in iter_jsonl(src):
            f.write((",\n" if n else "") + json.dumps(r, ensure_ascii=False))
            n += 1
        f.write("\n]\n")
    return n


LABELS_FA = {
    "name": "نام محصول", "url": "لینک", "price": "قیمت", "regular_price": "قیمت اصلی",
    "sale_price": "قیمت حراج", "min_price": "کمترین قیمت", "max_price": "بیشترین قیمت",
    "currency": "واحد پول", "in_stock": "موجودی", "categories": "دسته‌بندی", "tags": "برچسب‌ها",
    "brand": "برند", "sku": "کد کالا (SKU)", "product_id": "شناسه", "short_description": "توضیح کوتاه",
    "description": "توضیحات", "image_main": "عکس اصلی", "image_count": "تعداد عکس",
    "description_images": "عکس‌های داخل توضیحات", "files": "فایل‌ها (PDF و…)",
    "meta_title": "عنوان صفحه", "meta_description": "توضیح متا", "image_files": "فایل عکس‌های دانلودشده",
    "price_text": "متن قیمت", "source": "روش", "site": "سایت", "crawled_at": "زمان کراول",
    "description_html": "توضیحات (HTML)",
}
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _cell(v):
    if isinstance(v, bool):
        return "موجود" if v else "ناموجود"
    if isinstance(v, str):
        return _ILLEGAL.sub("", v)[:32000]
    return v


def _specs(r) -> dict:
    try:
        return json.loads(r["specs"]) if r.get("specs") else {}
    except ValueError:
        return {}


def make_thumbnail(path: str, size: int = 110) -> str | None:
    """PNG thumbnail next to the image (for embedding in Excel). None if Pillow can't read it."""
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        with Image.open(path) as im:
            im.thumbnail((size, size))
            if im.mode not in ("RGB", "RGBA"):
                im = im.convert("RGBA")
            out = os.path.splitext(path)[0] + "_thumb.png"
            im.save(out, "PNG", optimize=True)
            return out
    except Exception:
        return None


THUMB_MAX_ROWS = 1500        # above this the workbook is streamed (write-only) without embedded thumbnails
_HEAD_FIELDS = ["name", "url", "price", "regular_price", "sale_price", "min_price", "max_price",
                "currency", "in_stock", "categories", "tags", "brand", "sku", "product_id",
                "short_description", "description", "image_main", "image_count"]
_TAIL_FIELDS = ["description_images", "files", "meta_title", "meta_description", "image_files",
                "price_text", "source", "site", "crawled_at", "description_html"]


def _split(v) -> list[str]:
    return [u for u in (v or "").split(" | ") if u]


def write_xlsx(src, path: str, base_dir: str | None = None,
               max_image_cols: int = 10, max_spec_cols: int = 40) -> bool:
    """Excel with Persian headers: thumbnail (if images were downloaded), one column per image and per
    spec, plus 'specs' and 'images' sheets in long format. `src` is a .jsonl path (streamed in two
    passes) or a list of records. Large outputs are written in openpyxl's write-only mode."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    base_dir = base_dir or os.path.dirname(path)

    # pass 1: shape of the data only
    n_rows, n_img, has_files = 0, 0, False
    spec_counts: dict[str, int] = {}
    for r in iter_jsonl(src):
        n_rows += 1
        n_img = max(n_img, len(_split(r.get("images"))))
        has_files = has_files or bool(r.get("image_files"))
        for k in _specs(r):
            spec_counts[k] = spec_counts.get(k, 0) + 1
    n_img = min(n_img, max_image_cols)
    spec_keys = [k for k, _ in sorted(spec_counts.items(), key=lambda kv: -kv[1])][:max_spec_cols]
    streaming = n_rows > THUMB_MAX_ROWS
    has_thumbs = has_files and not streaming

    headers = (["تصویر"] if has_thumbs else []) + [LABELS_FA[f] for f in _HEAD_FIELDS] \
        + [f"عکس {i}" for i in range(1, n_img + 1)] + spec_keys + [LABELS_FA[f] for f in _TAIL_FIELDS]
    widths = {LABELS_FA["name"]: 40, LABELS_FA["url"]: 30, LABELS_FA["categories"]: 25,
              LABELS_FA["description"]: 60, LABELS_FA["short_description"]: 40, "تصویر": 17,
              LABELS_FA["description_html"]: 40, LABELS_FA["image_main"]: 30}
    wrap_cols = {headers.index(LABELS_FA["description"]), headers.index(LABELS_FA["short_description"])}
    first_img = headers.index("عکس 1") if n_img else -1
    link_cols = {headers.index(LABELS_FA["url"]), headers.index(LABELS_FA["image_main"])} \
        | (set(range(first_img, first_img + n_img)) if n_img else set())
    bold = Font(bold=True)
    mid = Alignment(vertical="center")
    mid_wrap = Alignment(vertical="center", wrap_text=True)

    wb = Workbook(write_only=streaming)
    if streaming:
        ws = wb.create_sheet("محصولات")
    else:
        ws = wb.active
        ws.title = "محصولات"
    ss = wb.create_sheet("مشخصات")
    si = wb.create_sheet("عکس‌ها")
    for sheet in (ws, ss, si):
        sheet.sheet_view.rightToLeft = True
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 22 if h.startswith("عکس ") else 15)
    for col, w in zip("ABCD", (40, 30, 25, 40)):
        ss.column_dimensions[col].width = w
    for col, w in zip("ABCDEF", (40, 30, 6, 14, 50, 30)):
        si.column_dimensions[col].width = w
    ws.freeze_panes = "C2" if has_thumbs else "B2"

    def header_row(sheet, values):
        if streaming:
            cells = []
            for v in values:
                c = WriteOnlyCell(sheet, value=v)
                c.font = bold
                cells.append(c)
            sheet.append(cells)
        else:
            sheet.append(values)
            for c in sheet[sheet.max_row]:
                c.font = bold

    header_row(ws, headers)
    header_row(ss, ["نام محصول", "لینک", "ویژگی", "مقدار"])
    header_row(si, ["نام محصول", "لینک محصول", "ردیف", "نوع", "لینک عکس", "فایل دانلودشده"])

    # pass 2: rows
    for ri, r in enumerate(iter_jsonl(src), start=2):
        imgs = _split(r.get("images"))
        specs = _specs(r)
        vals = ([""] if has_thumbs else []) + [_cell(r.get(f, "")) for f in _HEAD_FIELDS] \
            + [imgs[i] if i < len(imgs) else "" for i in range(n_img)] \
            + [_cell(specs.get(k, "")) for k in spec_keys] + [_cell(r.get(f, "")) for f in _TAIL_FIELDS]
        if streaming:
            row = []
            for ci, v in enumerate(vals):
                c = WriteOnlyCell(ws, value=v)
                c.alignment = mid_wrap if ci in wrap_cols else mid
                if ci in link_cols and isinstance(v, str) and v.startswith("http"):
                    c.hyperlink = v
                    c.style = "Hyperlink"
                row.append(c)
            ws.append(row)
        else:
            ws.append(vals)
            for ci, cell in enumerate(ws[ri]):
                cell.alignment = mid_wrap if ci in wrap_cols else mid
                if ci in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
            if has_thumbs:
                first_file = next(iter(_split(r.get("image_files"))), "")
                thumb = make_thumbnail(os.path.join(base_dir, first_file)) if first_file else None
                if thumb:
                    try:
                        pic = XLImage(thumb)
                        ws.add_image(pic, f"A{ri}")
                        ws.row_dimensions[ri].height = max(pic.height, 20) * 0.75 + 4
                    except Exception:
                        pass
        for k, v in specs.items():
            ss.append([_cell(r.get("name", "")), r.get("url", ""), _cell(k), _cell(v)])
        files = _split(r.get("image_files"))
        for i, u in enumerate(_dedupe(imgs + _split(r.get("description_images"))), 1):
            si.append([_cell(r.get("name", "")), r.get("url", ""), i,
                       "گالری" if i <= len(imgs) else "داخل توضیحات", u, files[i - 1] if i <= len(files) else ""])

    wb.save(path)
    return True


# --------------------------------------------------------------------------- #
# Site crawler
# --------------------------------------------------------------------------- #

def is_available(rec: dict) -> bool:
    """In stock (or stock unknown) AND has a real numeric price."""
    if rec.get("in_stock") is False:
        return False
    for k in ("price", "min_price", "sale_price", "regular_price"):
        v = rec.get(k)
        if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
            return True
        if isinstance(v, str) and v.strip():
            n = parse_number(v)
            if n and n > 0:
                return True
    return False


class SiteCrawler:
    def __init__(self, base: str, args, stop_event: threading.Event | None = None,
                 logger_name: str | None = None, out_path: str | None = None):
        if not base.endswith("/"):
            base += "/"
        self.base = base
        self.host = norm_host(urlparse(base).netloc)
        self.args = args
        self.stop_event = stop_event or threading.Event()
        self.fetcher = Fetcher(args, self.stop_event)
        self.robots = Robots("")
        self.log = logging.getLogger(logger_name or self.host)
        self.blocked_by_robots = 0
        self.only_available = bool(getattr(args, "only_available", False))
        self.skipped = 0
        self.skipped_urls: set[str] = set()
        self._skip_lock = threading.Lock()
        self._skip_path = (out_path or os.path.join(args.out, f"{self.host}.jsonl")) + ".skipped"
        if args.resume and os.path.exists(self._skip_path):
            with open(self._skip_path, encoding="utf-8") as f:
                self.skipped_urls = {line.strip() for line in f if line.strip()}
            self.skipped = len(self.skipped_urls)
        self._robots_cache: dict[str, Robots] = {}
        self._robots_lock = threading.Lock()
        self.writer = JsonlWriter(out_path or os.path.join(args.out, f"{self.host}.jsonl"), args.resume)

    @property
    def stopped(self) -> bool:
        return self.stop_event.is_set()

    # ---- plumbing
    def same_site(self, url: str) -> bool:
        """Same host, or another subdomain of the same domain (e.g. list.example.com for example.com)."""
        host = norm_host(urlparse(url).netloc)
        if host == self.host:
            return True
        return base_domain(host) == base_domain(self.host) and not host.startswith(SKIP_SUBDOMAINS)

    def robots_for(self, url: str) -> Robots:
        host = norm_host(urlparse(url).netloc)
        if host == self.host:
            return self.robots
        with self._robots_lock:
            if host not in self._robots_cache:
                r = self.fetcher.get(f"{urlparse(url).scheme}://{urlparse(url).netloc}/robots.txt")
                ok = r is not None and r.status_code == 200 and "text/html" not in r.headers.get("content-type", "")
                self._robots_cache[host] = Robots(r.text if ok else "")
            return self._robots_cache[host]

    def fetch(self, url: str):
        if not self.args.ignore_robots and not self.robots_for(url).allowed(url):
            self.blocked_by_robots += 1
            self.log.debug("robots.txt disallows %s", url)
            return None
        r = self.fetcher.get(url)
        if r is not None and r.status_code in (403, 503) and re.search(
                r"arvancloud|cloudflare|challenge|captcha|ddos", r.text[:5000], re.I):
            self.log.warning("blocked by firewall/challenge page (%s) at %s", r.status_code, url)
        return r

    def skip(self, rec: dict):
        url = rec.get("url") or ""
        with self._skip_lock:
            if url in self.skipped_urls:
                return
            self.skipped_urls.add(url)
            self.skipped += 1
            try:
                with open(self._skip_path, "a", encoding="utf-8") as f:
                    f.write(url + "\n")
            except OSError:
                pass
            if self.skipped % 100 == 0:
                self.log.info("skipped %d out-of-stock / no-price products", self.skipped)

    def keep(self, rec: dict | None) -> bool:
        """Applies the 'only in-stock with price' filter; records what was skipped."""
        if not rec or not rec.get("name"):
            return False
        if self.only_available and not is_available(rec):
            self.skip(rec)
            return False
        return True

    def emit(self, rec: dict | None) -> bool:
        if self.keep(rec):
            if self.writer.write(rec):
                if self.writer.count % 25 == 0:
                    self.log.info("%d products so far", self.writer.count)
                return True
        return False

    def load_robots(self):
        r = self.fetcher.get(urljoin(self.base, "/robots.txt"))
        if r is not None and r.status_code == 200 and "text/html" not in r.headers.get("content-type", ""):
            self.robots = Robots(r.text)
            if self.robots.crawl_delay and self.robots.crawl_delay > self.fetcher.delay:
                self.log.info("using crawl-delay %.1fs from robots.txt", self.robots.crawl_delay)
                self.fetcher.delay = self.robots.crawl_delay

    # ---- strategy 1: WooCommerce Store API
    def try_store_api(self) -> int:
        endpoints = [
            urljoin(self.base, "wp-json/wc/store/v1/products"),
            urljoin(self.base, "wp-json/wc/store/products"),
            urljoin(self.base, "?rest_route=/wc/store/v1/products"),
        ]
        for ep in endpoints:
            got, page, total_pages = 0, 1, None
            stock_filter = self.only_available
            while True:
                sep = "&" if "?" in ep else "?"
                params = {"per_page": 100, "page": page}
                if stock_filter:
                    params["stock_status"] = "instock"
                r = self.fetch(f"{ep}{sep}{urlencode(params)}")
                if r is not None and r.status_code == 400 and stock_filter and page == 1:
                    stock_filter = False             # older stores don't know the filter; filter locally
                    continue
                if r is None or r.status_code != 200:
                    break
                try:
                    data = r.json()
                except ValueError:
                    break
                if not isinstance(data, list) or not data:
                    break
                recs = [product_from_store_api(i, self.host) for i in data if isinstance(i, dict)]
                if self.only_available:
                    recs = [x for x in recs if self.keep(x)]     # don't open pages of products we'll drop
                if getattr(self.args, "enrich", False):
                    recs = self.enrich(recs)
                for rec in recs:
                    self.emit(rec)
                got += len(data)
                if self.stopped:
                    break
                if total_pages is None:
                    try:
                        total_pages = int(r.headers.get("X-WP-TotalPages") or 0)
                    except ValueError:
                        total_pages = 0
                    self.log.info("Store API works (%s pages)", total_pages or "?")
                if (total_pages and page >= total_pages) or len(data) < 100 or page >= 500:
                    break
                page += 1
            if got:
                return got
        return 0

    def enrich(self, recs: list[dict]) -> list[dict]:
        """Open each product page to add gallery images, spec tables, files and full content."""
        todo = [r for r in recs if r.get("url") and r["url"] not in self.writer.seen_urls]  # noqa
        if not todo:
            return recs
        self.log.info("reading %d product pages for full content/images", len(todo))
        pages = {}
        with ThreadPoolExecutor(max_workers=self.args.workers) as pool:
            for final_url, html in pool.map(lambda r: self._fetch_page(r["url"]), todo):
                if html:
                    try:
                        pages[final_url] = parse_product_page(html, final_url, self.host, trust_as_product=True)
                    except Exception as e:
                        self.log.debug("parse error %s: %s", final_url, e)
        out = []
        for r in recs:
            page = pages.get(r.get("url"))
            if page is None:            # follow redirects: match on path
                path = urlparse(r.get("url", "")).path.rstrip("/")
                page = next((p for u, p in pages.items() if urlparse(u).path.rstrip("/") == path), None)
            out.append(merge_records(r, page))
        return out

    # ---- images
    def download_images(self, mode: str = "main") -> int:
        """Download product images into <out>/images/, write 'image_files' back into the jsonl.
        mode: 'main' (first image only) or 'all' (gallery + description images, capped)."""
        if mode not in ("main", "all"):
            return 0
        if not os.path.exists(self.writer.path):
            return 0
        base_dir = os.path.dirname(self.writer.path)
        host_dir = re.sub(r"[^A-Za-z0-9.-]+", "_", self.host)
        max_bytes = int(float(os.environ.get("MAX_IMAGE_MB", "1500")) * 1024 * 1024)
        # images are static files: allow a shorter gap than for pages, but stay polite
        img_args = argparse.Namespace(**{**vars(self.args), "delay": max(0.1, self.args.delay / 4)})
        fetcher = Fetcher(img_args, self.stop_event)
        state = {"bytes": 0, "files": 0, "warned": False}
        lock = threading.Lock()

        def wanted(r):
            urls = [u for u in (r.get("images") or "").split(" | ") if u]
            if mode == "all":
                urls += [u for u in (r.get("description_images") or "").split(" | ") if u]
                return _dedupe(urls)[:MAX_IMAGES_PER_PRODUCT]
            return urls[:1]

        def one(idx_row):
            idx, r = idx_row
            folder = f"{idx + 1:04d}_{re.sub(r'[^A-Za-z0-9_-]', '', str(r.get('product_id') or r.get('sku') or ''))[:30]}".rstrip("_")
            saved = []
            for n, u in enumerate(wanted(r), 1):
                if self.stopped:
                    break
                with lock:
                    if state["bytes"] >= max_bytes:
                        if not state["warned"]:
                            self.log.warning("image size limit reached (%d MB); skipping the rest",
                                             max_bytes // 1048576)
                            state["warned"] = True
                        break
                if self.same_site(u) and not self.args.ignore_robots and not self.robots_for(u).allowed(u):
                    continue
                resp = fetcher.get(u)
                if resp is None or resp.status_code != 200:
                    continue
                ctype = resp.headers.get("content-type", "").split(";")[0].strip().lower()
                if not ctype.startswith("image/") or len(resp.content) > 15 * 1024 * 1024:
                    continue
                ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp", "image/gif": "gif",
                       "image/avif": "avif", "image/svg+xml": "svg"}.get(ctype, "img")
                rel = f"images/{host_dir}/{folder}/{n}.{ext}"
                os.makedirs(os.path.join(base_dir, "images", host_dir, folder), exist_ok=True)
                with open(os.path.join(base_dir, rel), "wb") as f:
                    f.write(resp.content)
                saved.append(rel)
                with lock:
                    state["bytes"] += len(resp.content)
                    state["files"] += 1
                    if state["files"] % 50 == 0:
                        self.log.info("downloaded %d images (%.0f MB)", state["files"], state["bytes"] / 1048576)
            r["image_files"] = " | ".join(saved)

        self.log.info("downloading images (%s) for %d products", mode, self.writer.count)
        tmp = self.writer.path + ".tmp"
        batch: list = []

        def flush(out):
            with ThreadPoolExecutor(max_workers=self.args.workers) as pool:
                list(pool.map(one, batch))
            for _, r in batch:
                out.write(json.dumps(r, ensure_ascii=False) + "\n")
            batch.clear()

        with open(tmp, "w", encoding="utf-8") as out:      # stream: only ~30 records in memory
            for idx, r in enumerate(iter_jsonl(self.writer.path)):
                batch.append((idx, r))
                if len(batch) >= 30:
                    flush(out)
            if batch:
                flush(out)
        os.replace(tmp, self.writer.path)
        self.log.info("images done: %d files, %.1f MB", state["files"], state["bytes"] / 1048576)
        return state["files"]

    # ---- strategy 2: sitemaps
    @staticmethod
    def _is_product_sitemap(u: str) -> bool:
        name = unquote(urlparse(u).path.rsplit("/", 1)[-1]).lower()
        if "product" not in name and "محصول" not in name:
            return False
        return not re.search(r"product[_-]?(cat|category|tag|brand)|pa_|taxonom", name)

    @staticmethod
    def _is_useless_sitemap(u: str) -> bool:
        name = unquote(urlparse(u).path.rsplit("/", 1)[-1]).lower()
        return bool(re.search(r"author|user|attachment|tag|portfolio|post_tag|local-sitemap|video|image|news", name))

    def _read_sitemap(self, url: str):
        r = self.fetch(url)
        if r is None or r.status_code != 200:
            return None, []
        content = r.content
        if content[:2] == b"\x1f\x8b":
            try:
                content = gzip.decompress(content)
            except OSError:
                return None, []
        try:
            root = ET.fromstring(content)
        except ET.ParseError:
            return None, []
        kind = root.tag.rsplit("}", 1)[-1].lower()
        locs = [el.text.strip() for el in root.iter()
                if el.tag.rsplit("}", 1)[-1].lower() == "loc" and el.text and el.text.strip()]
        return kind, locs

    def sitemap_urls(self) -> tuple[list[str], bool]:
        """Returns (urls, urls_are_known_products)."""
        declared = list(self.robots.sitemaps)
        fallbacks = [urljoin(self.base, p) for p in (
            "sitemap_index.xml", "product-sitemap.xml", "wp-sitemap.xml", "sitemap.xml",
            "sitemap-products.xml", "sitemap/products.xml")]
        product_urls, other_urls, seen = set(), set(), set()

        def walk(start: str):
            q = deque([start])
            while q and len(seen) < 1000:
                sm = q.popleft()
                if sm in seen:
                    continue
                seen.add(sm)
                kind, locs = self._read_sitemap(sm)
                if kind is None:
                    continue
                if kind == "sitemapindex":
                    prod = [l for l in locs if self._is_product_sitemap(l)]
                    q.extend(prod or [l for l in locs if not self._is_useless_sitemap(l)])
                elif self._is_product_sitemap(sm):
                    product_urls.update(locs)
                else:
                    other_urls.update(locs)

        for sm in declared:
            walk(sm)
        for sm in fallbacks:
            if product_urls:
                break
            walk(sm)

        keep = lambda u: self.same_site(u) and not is_excluded(u)
        if product_urls:
            return sorted(u for u in product_urls if keep(u)), True
        others = [u for u in other_urls if keep(u)]
        guessed = [u for u in others if looks_like_product_url(u)]
        if guessed:
            return sorted(guessed), False
        return sorted(others), False

    # ---- page workers
    def _fetch_page(self, url: str):
        r = self.fetch(url)
        if r is None or r.status_code != 200:
            return url, None
        ctype = r.headers.get("content-type", "")
        if "html" not in ctype and ctype:
            return url, None
        r.encoding = r.encoding if r.encoding and r.encoding.lower() != "iso-8859-1" else "utf-8"
        return r.url, r.text

    def fetch_stream(self, pool, next_url, window: int):
        """Yield (final_url, html) as downloads finish, keeping `window` requests in flight so the network
        stays busy while pages are parsed. Memory is bounded by the window, not by the number of URLs."""
        inflight = set()

        def fill():
            while len(inflight) < window and not self.stopped:
                u = next_url()
                if u is None:
                    break
                inflight.add(pool.submit(self._fetch_page, u))

        fill()
        while inflight:
            done, _ = wait(inflight, return_when=FIRST_COMPLETED)
            for f in done:
                inflight.discard(f)
                yield f.result()
            fill()

    def crawl_pages(self, urls: list[str], trust: bool) -> int:
        todo = [u for u in urls if u not in self.writer.seen_urls and u not in self.skipped_urls][: self.args.max_pages]
        self.log.info("parsing %d candidate pages from sitemap", len(todo))
        found, done = 0, 0
        it = iter(todo)
        with ThreadPoolExecutor(max_workers=self.args.workers) as pool:
            for final_url, html in self.fetch_stream(pool, lambda: next(it, None), self.args.workers * 2):
                done += 1
                if html:
                    found += self.parse_any(html, final_url, trust)
                    del html
                if done % 200 == 0:
                    self.log.info("pages %d/%d, %d products", done, len(todo), self.writer.count)
        return found

    def parse_any(self, html: str, url: str, trust: bool = False) -> int:
        """A product page -> one record; otherwise look for a price table / embedded product list."""
        try:
            rec = parse_product_page(html, url, self.host, trust_as_product=trust)
        except Exception as e:
            self.log.debug("parse error %s: %s", url, e)
            rec = None
        if rec:
            return int(self.emit(rec))
        if "<table" not in html and "application/json" not in html:
            return 0                       # nothing a price list could be made of; skip the second parse
        try:
            recs = extract_listing_records(html, url, self.host,
                                           require_price=getattr(self.args, "mode", "") != "list")
        except Exception as e:
            self.log.debug("listing parse error %s: %s", url, e)
            return 0
        n = sum(self.emit(r) for r in recs)
        if n:
            self.log.info("price list on %s: %d items", url, n)
        return n

    # ---- strategy 3: BFS link crawl
    def extract_links(self, html: str, page_url: str) -> list[str]:
        try:
            root = lxml.html.fromstring(html.encode("utf-8", "replace"), parser=_LXML_PARSER)
        except Exception:
            return []
        out = []
        for a in root.iter("a"):
            href = (a.get("href") or "").strip()
            if not href or href.startswith(("mailto:", "tel:", "javascript:", "#", "whatsapp:")):
                continue
            u = normalize_url(urljoin(page_url, href))
            if urlparse(u).scheme in ("http", "https") and self.same_site(u) and not is_excluded(u):
                out.append(u)
        return out

    def bfs(self) -> int:
        starts = [self.base] + [urljoin(self.base, p) for p in ("shop/", "products/", "store/", "فروشگاه/")]
        hi, lo = deque(), deque()
        seen = set()
        for s in starts:
            s = normalize_url(s)
            seen.add(s)
            hi.append(s)
        state = {"submitted": 0, "pages": 0}
        found = 0

        def next_url():
            if state["submitted"] >= self.args.max_pages or not (hi or lo):
                return None
            state["submitted"] += 1
            return hi.popleft() if hi else lo.popleft()

        with ThreadPoolExecutor(max_workers=self.args.workers) as pool:
            for final_url, html in self.fetch_stream(pool, next_url, self.args.workers * 2):
                state["pages"] += 1
                if html:
                    if final_url not in self.writer.seen_urls and final_url not in self.skipped_urls:
                        found += self.parse_any(html, final_url)
                    for link in self.extract_links(html, final_url):
                        if link in seen:
                            continue
                        seen.add(link)
                        promising = looks_like_product_url(link) or looks_like_listing_url(link) or \
                            re.search(r"(^|\.)(list|price|prices|shop|store|catalog|pricelist)\.", urlparse(link).netloc)
                        (hi if promising else lo).append(link)
                    del html
                if state["pages"] % 100 == 0:
                    self.log.info("crawled %d pages, %d queued, %d products", state["pages"], len(hi) + len(lo),
                                  self.writer.count)
        return found

    # ---- main
    def run(self) -> int:
        t0 = time.time()
        self.log.info("start")
        try:
            self.load_robots()
            home = self.fetcher.get(self.base)
            if home is None and self.base.startswith("https://") and not self.stopped:
                self.log.warning("https failed, trying http://")
                self.base = "http://" + self.base[len("https://"):]
                self.load_robots()
                home = self.fetcher.get(self.base)
            if home is None:
                self.log.error("site unreachable (network / DNS / geo-block?). Try --proxy.")
                return 0
            if home.status_code >= 400:
                self.log.warning("home page returned HTTP %s", home.status_code)

            mode, total = self.args.mode, 0
            if mode in ("auto", "api"):
                total = self.try_store_api()
                if total:
                    self.log.info("Store API: %d products", total)
            if not total and mode in ("auto", "sitemap") and not self.stopped:
                urls, trusted = self.sitemap_urls()
                if urls:
                    self.log.info("sitemap: %d urls (%s)", len(urls), "product sitemap" if trusted else "guessed")
                    total = self.crawl_pages(urls, trusted)
            if not total and mode in ("auto", "crawl", "list") and not self.stopped:
                self.log.info("falling back to link crawling (max %d pages)", self.args.max_pages)
                total = self.bfs()
            if self.blocked_by_robots:
                self.log.warning("%d URLs skipped because robots.txt disallows them "
                                 "(use --ignore-robots only if you have the owner's permission)",
                                 self.blocked_by_robots)
            if self.stopped:
                self.log.warning("stopped (توقف شد)")
            if self.only_available:
                self.log.info("skipped %d products (out of stock or no price)", self.skipped)
            self.log.info("done: %d products in %.0fs", self.writer.count, time.time() - t0)
            return self.writer.count
        finally:
            self.writer.close()


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser(description="Crawl products from WooCommerce / generic shop sites.")
    ap.add_argument("--sites", nargs="*", help="site URLs (default: the built-in list)")
    ap.add_argument("--sites-file", help="text file with one site URL per line")
    ap.add_argument("--out", default="output", help="output folder (default: output)")
    ap.add_argument("--mode", choices=["auto", "api", "sitemap", "crawl", "list"], default="auto")
    ap.add_argument("--delay", type=float, default=1.0, help="seconds between requests per site (default 1.0)")
    ap.add_argument("--workers", type=int, default=3, help="parallel requests per site (default 3)")
    ap.add_argument("--parallel-sites", type=int, default=4, help="sites crawled at the same time (default 4)")
    ap.add_argument("--max-pages", type=int, default=5000, help="max HTML pages per site (default 5000)")
    ap.add_argument("--timeout", type=float, default=30)
    ap.add_argument("--proxy", help="http(s)/socks proxy, e.g. http://127.0.0.1:8080")
    ap.add_argument("--user-agent", default=DEFAULT_USER_AGENT)
    ap.add_argument("--images", choices=["none", "main", "all"], default="none",
                    help="download product images: none / main (first image) / all")
    ap.add_argument("--enrich", action="store_true",
                    help="when the Store API works, still open each product page for full content & gallery")
    ap.add_argument("--only-available", action="store_true",
                    help="keep only in-stock products that have a price")
    ap.add_argument("--resume", action="store_true", help="keep existing .jsonl and skip already-saved URLs")
    ap.add_argument("--ignore-robots", action="store_true",
                    help="do not obey robots.txt (only with the site owner's permission)")
    ap.add_argument("--insecure", action="store_true", help="skip TLS certificate verification")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(name)-22s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")
    logging.getLogger("urllib3").setLevel(logging.WARNING)

    sites = list(args.sites or [])
    if args.sites_file:
        with open(args.sites_file, encoding="utf-8") as f:
            sites += [l.strip() for l in f if l.strip() and not l.startswith("#")]
    if not sites:
        sites = DEFAULT_SITES
    sites = _dedupe(s if s.startswith("http") else "https://" + s for s in sites)

    os.makedirs(args.out, exist_ok=True)
    crawlers = [SiteCrawler(s, args) for s in sites]
    summary = {}
    with ThreadPoolExecutor(max_workers=max(1, args.parallel_sites)) as pool:
        def run_site(c):
            n = c.run()
            if args.images != "none" and n:
                c.download_images(args.images)
            return n
        futs = {pool.submit(run_site, c): c for c in crawlers}
        for fut in as_completed(futs):
            c = futs[fut]
            try:
                summary[c.host] = fut.result()
            except Exception as e:
                logging.getLogger(c.host).exception("crashed: %s", e)
                summary[c.host] = 0

    combined = os.path.join(args.out, "all_products.jsonl")
    total = 0
    with open(combined, "w", encoding="utf-8") as out:
        for c in crawlers:
            if os.path.exists(c.writer.path):
                write_csv(c.writer.path, os.path.join(args.out, f"{c.host}.csv"))
                with open(c.writer.path, encoding="utf-8") as f:
                    for line in f:
                        out.write(line)
                        total += 1
    write_csv(combined, os.path.join(args.out, "all_products.csv"))
    xlsx = write_xlsx(combined, os.path.join(args.out, "all_products.xlsx"))

    print("\n===== summary =====")
    for c in crawlers:
        print(f"  {c.host:<28} {summary.get(c.host, 0):>6} products")
    print(f"  {'TOTAL':<28} {total:>6}")
    print(f"\nfiles in: {os.path.abspath(args.out)}" + ("" if xlsx else "  (pip install openpyxl for .xlsx)"))


if __name__ == "__main__":
    main()
