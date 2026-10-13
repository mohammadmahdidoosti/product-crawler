#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Product crawler for WordPress / WooCommerce (and generic) shop sites.

For every site it tries, in order:
  1. WooCommerce Store API  (/wp-json/wc/store/v1/products)  -> fastest, cleanest data
  2. XML sitemaps            (product-sitemap.xml, wp-sitemap-posts-product-*.xml, ...)
  3. Link crawling (BFS)     starting from the home page / shop page
Product pages are parsed from JSON-LD (schema.org Product), OpenGraph and
WooCommerce HTML markup, so it also works on non-WooCommerce sites that expose
structured data.

Output (in ./output by default):
  <domain>.jsonl      one product per line, written as it is crawled (resumable)
  <domain>.csv        per-site table (UTF-8 with BOM, opens correctly in Excel)
  all_products.csv    every site combined
  all_products.xlsx   every site combined (if openpyxl is installed)

Usage:
  python crawler.py                         # crawl all default sites
  python crawler.py --sites https://tsaco-diesel.com/ https://parsshoa.co/
  python crawler.py --mode sitemap --delay 2 --max-pages 3000 --resume
  python crawler.py --proxy http://127.0.0.1:8080
"""

from __future__ import annotations

import argparse
import csv
import gzip
import html as html_lib
import json
import logging
import os
import random
import re
import sys
import threading
import time
import xml.etree.ElementTree as ET
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from urllib.parse import parse_qsl, unquote, urldefrag, urlencode, urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

DEFAULT_SITES = [
    "https://tsaco-diesel.com/",
    "https://parsshoa.co/",
    "https://foolad-parsekaveh.com/",
    "https://mahyaelectronic.ir/",
    "https://yahyabrand.com/",
    "https://alvandlift.ir/",
    "https://safetypatan.com/",
    "https://saterco.ir/",
]

DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/129.0 Safari/537.36 ProductCrawler/1.0"
)
ROBOTS_AGENT_TOKEN = "ProductCrawler"

FIELDS = [
    "site", "source", "url", "product_id", "name", "sku", "price",
    "regular_price", "sale_price", "min_price", "max_price", "currency",
    "price_text", "in_stock", "categories", "brand", "short_description",
    "description", "attributes", "images", "crawled_at",
]

# URL path fragments that usually indicate a single product page
PRODUCT_URL_PATTERNS = [
    r"/product/", r"/products/[^/]+", r"/shop/[^/]+/?$", r"/%D9%85%D8%AD%D8%B5%D9%88%D9%84/",
    r"/محصول/", r"/item/", r"/p/\d+", r"/kala/", r"/کالا/",
]
# URL path fragments that are never products and not worth crawling
EXCLUDE_URL_PATTERNS = [
    r"/wp-admin", r"/wp-login", r"/wp-json", r"/wp-content/", r"/wp-includes/",
    r"/cart", r"/checkout", r"/my-account", r"/account", r"/login", r"/register",
    r"/feed", r"/comments/", r"/author/", r"/xmlrpc", r"/wishlist", r"/compare",
    r"/سبد-خرید", r"/تسویه-حساب", r"/حساب-کاربری",
    r"add-to-cart=", r"add_to_wishlist", r"replytocom=", r"orderby=", r"filter_",
    r"[?&]s=", r"[?&]min_price=", r"[?&]max_price=", r"/tag/", r"/product-tag/",
    r"\.(?:jpe?g|png|gif|webp|svg|pdf|zip|rar|mp4|mp3|avi|docx?|xlsx?|css|js|ico|woff2?|ttf)(?:$|\?)",
]
# Query keys kept when normalising URLs (pagination); everything else is dropped
KEEP_QUERY_KEYS = {"page", "paged", "product-page", "p", "pg"}

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

log = logging.getLogger("crawler")


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #

def to_latin_digits(s: str) -> str:
    return (s or "").translate(PERSIAN_DIGITS)


def clean_text(s) -> str:
    if s is None:
        return ""
    if not isinstance(s, str):
        s = str(s)
    if "<" in s and ">" in s:
        s = BeautifulSoup(s, "lxml").get_text("\n")
    s = html_lib.unescape(s).replace("\xa0", " ")
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n", s)
    return s.strip()


def parse_number(token: str):
    """'1,800,000' / '۱٬۸۰۰٬۰۰۰' / '11.000.000.000' / '12.5' -> float"""
    t = to_latin_digits(token).strip()
    t = t.replace("٬", ",").replace("،", ",").replace(" ", "")
    if not t:
        return None
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", t):           # dots as thousands sep
        t = t.replace(".", "")
    t = t.replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def extract_numbers(text: str) -> list[float]:
    t = to_latin_digits(text or "")
    out = []
    for tok in re.findall(r"\d[\d,٬،.]*\d|\d", t):
        n = parse_number(tok)
        if n is not None:
            out.append(n)
    return out


def detect_currency(text: str) -> str:
    t = text or ""
    if "تومان" in t or "toman" in t.lower():
        return "IRT"
    if "ریال" in t or "﷼" in t or "rial" in t.lower():
        return "IRR"
    return ""


def norm_host(host: str) -> str:
    host = (host or "").lower()
    return host[4:] if host.startswith("www.") else host


def normalize_url(url: str) -> str:
    url, _ = urldefrag(url)
    p = urlparse(url)
    q = [(k, v) for k, v in parse_qsl(p.query, keep_blank_values=True) if k in KEEP_QUERY_KEYS]
    path = p.path or "/"
    return p._replace(path=path, query=urlencode(q), fragment="").geturl()


def is_excluded(url: str) -> bool:
    u = unquote(url).lower()
    return any(re.search(pat, u) for pat in EXCLUDE_URL_PATTERNS)


def looks_like_product_url(url: str) -> bool:
    path = urlparse(url).path
    if re.search(r"/product-(?:cat|category|tag|brand)/", path, re.I):
        return False
    return any(re.search(p, path, re.I) or re.search(p, unquote(path), re.I) for p in PRODUCT_URL_PATTERNS)


def looks_like_listing_url(url: str) -> bool:
    u = unquote(url).lower()
    return bool(re.search(r"/(shop|store|products?|product-cat|product-category|category|"
                          r"فروشگاه|محصولات|دسته)|/page/\d+|[?&](page|paged|product-page)=", u))


# --------------------------------------------------------------------------- #
# robots.txt (supports * and $ wildcards, longest-match rule)
# --------------------------------------------------------------------------- #

class Robots:
    def __init__(self, text: str = "", agent: str = ROBOTS_AGENT_TOKEN):
        self.rules: list[tuple[bool, str]] = []
        self.crawl_delay: float | None = None
        self.sitemaps: list[str] = []
        self._parse(text or "", agent.lower())

    def _parse(self, text: str, agent: str):
        groups = []           # (agents, rules, delay)
        agents, rules, delay = [], [], None
        last_was_agent = False
        for raw in text.splitlines():
            line = raw.split("#", 1)[0].strip()
            if ":" not in line:
                continue
            key, val = line.split(":", 1)
            key, val = key.strip().lower(), val.strip()
            if key == "user-agent":
                if not last_was_agent and agents:
                    groups.append((agents, rules, delay))
                    agents, rules, delay = [], [], None
                agents.append(val.lower())
                last_was_agent = True
                continue
            last_was_agent = False
            if key in ("allow", "disallow"):
                rules.append((key == "allow", val))
            elif key == "crawl-delay":
                try:
                    delay = float(val)
                except ValueError:
                    pass
            elif key == "sitemap":
                self.sitemaps.append(val)
        if agents:
            groups.append((agents, rules, delay))

        specific = [g for g in groups if any(a != "*" and a in agent for a in g[0])]
        chosen = specific or [g for g in groups if "*" in g[0]]
        for _, r, d in chosen:
            self.rules.extend(r)
            if d is not None:
                self.crawl_delay = d

    @staticmethod
    def _match_len(pattern: str, path: str) -> int:
        pattern = unquote(pattern)
        anchored = pattern.endswith("$")
        if anchored:
            pattern = pattern[:-1]
        regex = "^" + "".join(".*" if c == "*" else re.escape(c) for c in pattern)
        if anchored:
            regex += "$"
        return len(pattern) if re.match(regex, path) else -1

    def allowed(self, url: str) -> bool:
        p = urlparse(url)
        path = unquote(p.path or "/") + (("?" + unquote(p.query)) if p.query else "")
        best_len, best_allow = -1, True
        for allow, pat in self.rules:
            if not pat:
                continue
            n = self._match_len(pat, path)
            if n > best_len or (n == best_len and allow):
                best_len, best_allow = n, allow
        return best_allow if best_len >= 0 else True


# --------------------------------------------------------------------------- #
# HTTP fetching with retries + per-site rate limit
# --------------------------------------------------------------------------- #

class Fetcher:
    def __init__(self, args, stop_event: threading.Event | None = None):
        self.stop_event = stop_event
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": args.user_agent,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
            "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.7",
        })
        retry = Retry(total=3, connect=3, read=3, backoff_factor=1.5,
                      status_forcelist=(429, 500, 502, 503, 504),
                      allowed_methods=("GET", "HEAD"), respect_retry_after_header=True)
        adapter = HTTPAdapter(max_retries=retry, pool_connections=10, pool_maxsize=10)
        self.session.mount("http://", adapter)
        self.session.mount("https://", adapter)
        if args.proxy:
            self.session.proxies.update({"http": args.proxy, "https": args.proxy})
        self.session.verify = not args.insecure
        self.timeout = args.timeout
        self.delay = args.delay
        self._lock = threading.Lock()
        self._next = 0.0

    def get(self, url: str):
        if self.stop_event is not None and self.stop_event.is_set():
            return None
        with self._lock:                          # polite spacing between requests
            wait = self._next - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            self._next = time.monotonic() + self.delay * random.uniform(0.8, 1.3)
        try:
            return self.session.get(url, timeout=self.timeout, allow_redirects=True)
        except requests.RequestException as e:
            log.debug("request failed %s: %s", url, e)
            return None


# --------------------------------------------------------------------------- #
# Product page parsing
# --------------------------------------------------------------------------- #

def _iter_jsonld(soup: BeautifulSoup):
    for tag in soup.find_all("script", attrs={"type": re.compile(r"ld\+json", re.I)}):
        raw = tag.string or tag.get_text() or ""
        raw = re.sub(r"[\x00-\x1f]", " ", raw).strip()
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except ValueError:
            continue
        stack = [data]
        while stack:
            node = stack.pop()
            if isinstance(node, list):
                stack.extend(node)
            elif isinstance(node, dict):
                yield node
                for k in ("@graph", "mainEntity", "itemListElement", "item"):
                    if k in node:
                        stack.append(node[k])


def _is_type(node: dict, name: str) -> bool:
    t = node.get("@type")
    types = t if isinstance(t, list) else [t]
    return any(isinstance(x, str) and x.lower() == name.lower() for x in types)


def find_jsonld_product(soup: BeautifulSoup) -> dict | None:
    best = None
    for node in _iter_jsonld(soup):
        if _is_type(node, "Product") or _is_type(node, "ProductGroup"):
            if best is None or (node.get("offers") and not best.get("offers")):
                best = node
    return best


def _as_list(x):
    if x is None:
        return []
    return x if isinstance(x, list) else [x]


def _image_urls(x) -> list[str]:
    out = []
    for item in _as_list(x):
        if isinstance(item, str):
            out.append(item)
        elif isinstance(item, dict):
            u = item.get("url") or item.get("contentUrl") or item.get("@id")
            if isinstance(u, str):
                out.append(u)
    return out


def _dedupe(seq):
    seen, out = set(), []
    for x in seq:
        if x and x not in seen:
            seen.add(x)
            out.append(x)
    return out


def first(soup, selectors: list[str]):
    """First element matching the selectors, tried in priority order."""
    for sel in selectors:
        el = soup.select_one(sel)
        if el is not None:
            return el
    return None


def parse_product_page(html: str, url: str, site: str, trust_as_product: bool = False) -> dict | None:
    soup = BeautifulSoup(html, "lxml")
    ld = find_jsonld_product(soup)
    meta = {}
    for m in soup.find_all("meta"):
        k = (m.get("property") or m.get("name") or m.get("itemprop") or "").strip().lower()
        if k and m.get("content") and k not in meta:
            meta[k] = m["content"].strip()
    body_cls = " ".join((soup.body.get("class") or [])) if soup.body else ""

    html_signals = bool(
        soup.select_one("form.cart, .product_title, .single-product-summary, "
                        "[itemtype*='schema.org/Product'], .woocommerce-product-gallery")
    )
    is_product = (
        ld is not None
        or meta.get("og:type", "").lower() == "product"
        or "single-product" in body_cls
        or html_signals
    )
    if not is_product and not trust_as_product:
        return None

    rec = {k: "" for k in FIELDS}
    rec.update(site=site, source="page", url=url,
               crawled_at=datetime.now(timezone.utc).isoformat(timespec="seconds"))

    # ---- name
    h1 = first(soup, ["h1.product_title", "h1.product-title", ".product_title", "main h1", "h1"])
    rec["name"] = clean_text((ld or {}).get("name")) or (clean_text(h1.get_text()) if h1 else "") \
        or clean_text(meta.get("og:title", ""))
    if not rec["name"]:
        return None

    # ---- id / sku
    rec["sku"] = clean_text((ld or {}).get("sku") or (ld or {}).get("mpn") or "")
    if not rec["sku"]:
        el = soup.select_one(".sku, [itemprop=sku]")
        if el:
            rec["sku"] = clean_text(el.get("content") or el.get_text())
    m = re.search(r"postid-(\d+)", body_cls)
    if m:
        rec["product_id"] = m.group(1)
    else:
        el = soup.select_one("[data-product_id], input[name=add-to-cart], button[name=add-to-cart]")
        if el:
            rec["product_id"] = el.get("data-product_id") or el.get("value") or ""

    # ---- brand
    brand = (ld or {}).get("brand")
    if isinstance(brand, list):
        brand = brand[0] if brand else ""
    if isinstance(brand, dict):
        brand = brand.get("name", "")
    rec["brand"] = clean_text(brand)

    # ---- price (JSON-LD offers first)
    prices, currency, availability = [], "", ""
    for off in _as_list((ld or {}).get("offers")):
        if not isinstance(off, dict):
            continue
        currency = currency or off.get("priceCurrency", "")
        availability = availability or str(off.get("availability", ""))
        for key in ("price", "lowPrice", "highPrice"):
            v = off.get(key)
            if v not in (None, ""):
                n = parse_number(str(v))
                if n is not None:
                    prices.append(n)
        spec = off.get("priceSpecification")
        for ps in _as_list(spec):
            if isinstance(ps, dict) and ps.get("price") not in (None, ""):
                n = parse_number(str(ps["price"]))
                if n is not None:
                    prices.append(n)
                currency = currency or ps.get("priceCurrency", "")

    price_el = first(soup, [".summary .price", ".entry-summary .price", ".product-info .price",
                            "p.price", "[itemprop=price]", ".product-price", "span.price"])
    price_text = ""
    if price_el:
        price_text = clean_text(price_el.get("content") or price_el.get_text(" "))
        ins, dels = price_el.select_one("ins"), price_el.select_one("del")
        if ins and dels:
            sale = extract_numbers(ins.get_text(" "))
            reg = extract_numbers(dels.get_text(" "))
            if sale:
                rec["sale_price"] = sale[0]
            if reg:
                rec["regular_price"] = reg[0]
    if not price_text and meta.get("product:price:amount"):
        price_text = meta["product:price:amount"]
    rec["price_text"] = price_text

    if not prices and price_text:
        nums = [n for n in extract_numbers(price_text) if n > 0]
        # with sale markup the "current" price is the <ins> one
        prices = [rec["sale_price"]] if rec["sale_price"] != "" else nums
    if prices:
        rec["min_price"], rec["max_price"] = min(prices), max(prices)
        rec["price"] = rec["sale_price"] if rec["sale_price"] != "" else min(prices)
    rec["currency"] = currency or meta.get("product:price:currency", "") or detect_currency(price_text)

    # ---- stock
    av = availability.lower()
    stock_el = soup.select_one(".summary .stock, p.stock, .stock")
    stock_txt = clean_text(stock_el.get_text()) if stock_el else ""
    stock_cls = " ".join(stock_el.get("class") or []) if stock_el else ""
    if "outofstock" in av or "out-of-stock" in stock_cls or "ناموجود" in stock_txt:
        rec["in_stock"] = False
    elif "instock" in av or "in-stock" in stock_cls or "موجود" in stock_txt:
        rec["in_stock"] = True

    # ---- categories
    cats = [clean_text(a.get_text()) for a in soup.select(".product_meta .posted_in a, .posted_in a")]
    if not cats:
        crumbs = soup.select(".woocommerce-breadcrumb a, nav.rank-math-breadcrumb a, "
                             ".breadcrumb a, .breadcrumbs a, [class*=breadcrumb] a")
        cats = [clean_text(a.get_text()) for a in crumbs[1:]]
    if not cats and ld and ld.get("category"):
        cats = [clean_text(c) for c in _as_list(ld["category"]) if isinstance(c, str)]
    rec["categories"] = " | ".join(_dedupe([c for c in cats if c and c != rec["name"]]))

    # ---- descriptions
    sd = first(soup, [".woocommerce-product-details__short-description", ".short-description",
                      "div[itemprop=description]"])
    rec["short_description"] = clean_text(sd.get_text("\n")) if sd else ""
    desc_el = first(soup, ["#tab-description", ".woocommerce-Tabs-panel--description",
                           ".product-description", "#description"])
    rec["description"] = (clean_text(desc_el.get_text("\n")) if desc_el else "") \
        or clean_text((ld or {}).get("description")) \
        or rec["short_description"] or clean_text(meta.get("description", ""))

    # ---- attributes / specs
    attrs = {}
    for tr in soup.select("table.woocommerce-product-attributes tr, table.shop_attributes tr"):
        th, td = tr.find("th"), tr.find("td")
        if th and td:
            attrs[clean_text(th.get_text())] = clean_text(td.get_text(" "))
    for prop in _as_list((ld or {}).get("additionalProperty")):
        if isinstance(prop, dict) and prop.get("name"):
            attrs.setdefault(clean_text(prop["name"]), clean_text(prop.get("value")))
    rec["attributes"] = json.dumps(attrs, ensure_ascii=False) if attrs else ""

    # ---- images
    imgs = _image_urls((ld or {}).get("image"))
    for a in soup.select(".woocommerce-product-gallery__image a[href], .woocommerce-product-gallery a[href]"):
        imgs.append(a["href"])
    for img in soup.select(".woocommerce-product-gallery img"):
        u = img.get("data-large_image") or img.get("data-src") or img.get("src")
        if u and not u.startswith("data:") and "lazy" not in u:
            imgs.append(u)
    if meta.get("og:image"):
        imgs.append(meta["og:image"])
    rec["images"] = " | ".join(_dedupe(urljoin(url, u) for u in imgs if isinstance(u, str)))
    return rec


def product_from_store_api(item: dict, site: str) -> dict:
    rec = {k: "" for k in FIELDS}
    prices = item.get("prices") or {}
    try:
        mu = int(prices.get("currency_minor_unit") or 0)
    except (TypeError, ValueError):
        mu = 0

    def money(v):
        if v in (None, ""):
            return ""
        try:
            return int(v) / (10 ** mu) if mu else int(v)
        except (TypeError, ValueError):
            n = parse_number(str(v))
            return n if n is not None else ""

    price, reg, sale = money(prices.get("price")), money(prices.get("regular_price")), money(prices.get("sale_price"))
    rng = prices.get("price_range") or {}
    mn, mx = money(rng.get("min_amount")), money(rng.get("max_amount"))
    if sale != "" and sale == reg:
        sale = ""

    attrs = {}
    for a in item.get("attributes") or []:
        terms = [t.get("name", "") for t in (a.get("terms") or []) if isinstance(t, dict)]
        attrs[clean_text(a.get("name"))] = "، ".join(clean_text(t) for t in terms)

    brand = ""
    for b in item.get("brands") or []:
        if isinstance(b, dict) and b.get("name"):
            brand = clean_text(b["name"])
            break

    rec.update(
        site=site, source="store_api",
        url=item.get("permalink", ""),
        product_id=item.get("id", ""),
        name=clean_text(item.get("name")),
        sku=clean_text(item.get("sku")),
        price=price, regular_price=reg, sale_price=sale,
        min_price=mn if mn != "" else price, max_price=mx if mx != "" else price,
        currency=prices.get("currency_code") or prices.get("currency_symbol") or "",
        price_text=clean_text(item.get("price_html")),
        in_stock=item.get("is_in_stock", ""),
        categories=" | ".join(clean_text(c.get("name")) for c in item.get("categories") or [] if isinstance(c, dict)),
        brand=brand,
        short_description=clean_text(item.get("short_description")),
        description=clean_text(item.get("description")),
        attributes=json.dumps(attrs, ensure_ascii=False) if attrs else "",
        images=" | ".join(_dedupe(i.get("src", "") for i in item.get("images") or [] if isinstance(i, dict))),
        crawled_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
    )
    return rec


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #

class JsonlWriter:
    def __init__(self, path: str, resume: bool):
        self.path = path
        self.lock = threading.Lock()
        self.seen_urls: set[str] = set()
        self.count = 0
        if resume and os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                for line in f:
                    try:
                        self.seen_urls.add(json.loads(line)["url"])
                        self.count += 1
                    except (ValueError, KeyError):
                        pass
            self.fh = open(path, "a", encoding="utf-8")
        else:
            self.fh = open(path, "w", encoding="utf-8")

    def write(self, rec: dict) -> bool:
        key = rec.get("url") or f"id:{rec.get('product_id')}"
        with self.lock:
            if key in self.seen_urls:
                return False
            self.seen_urls.add(key)
            self.fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
            self.fh.flush()
            self.count += 1
            return True

    def close(self):
        self.fh.close()


def read_jsonl(path: str) -> list[dict]:
    rows = []
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            for line in f:
                try:
                    rows.append(json.loads(line))
                except ValueError:
                    pass
    return rows


def write_csv(rows: list[dict], path: str):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in FIELDS})


def write_xlsx(rows: list[dict], path: str) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.sheet_view.rightToLeft = True
    ws.append(FIELDS)
    illegal = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
    for r in rows:
        vals = []
        for k in FIELDS:
            v = r.get(k, "")
            if isinstance(v, str):
                v = illegal.sub("", v)[:32000]
            vals.append(v)
        ws.append(vals)
    widths = {"name": 45, "url": 50, "categories": 30, "description": 60, "short_description": 40, "images": 40}
    for i, k in enumerate(FIELDS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 14)
    ws.freeze_panes = "A2"
    wb.save(path)
    return True


# --------------------------------------------------------------------------- #
# Site crawler
# --------------------------------------------------------------------------- #

class SiteCrawler:
    def __init__(self, base: str, args, stop_event: threading.Event | None = None,
                 logger_name: str | None = None, out_path: str | None = None):
        if not base.endswith("/"):
            base += "/"
        self.base = base
        self.host = norm_host(urlparse(base).netloc)
        self.args = args
        self.stop_event = stop_event or threading.Event()
        self.fetcher = Fetcher(args, self.stop_event)
        self.robots = Robots("")
        self.log = logging.getLogger(logger_name or self.host)
        self.blocked_by_robots = 0
        self.writer = JsonlWriter(out_path or os.path.join(args.out, f"{self.host}.jsonl"), args.resume)

    @property
    def stopped(self) -> bool:
        return self.stop_event.is_set()

    # ---- plumbing
    def same_site(self, url: str) -> bool:
        return norm_host(urlparse(url).netloc) == self.host

    def fetch(self, url: str):
        if not self.args.ignore_robots and not self.robots.allowed(url):
            self.blocked_by_robots += 1
            self.log.debug("robots.txt disallows %s", url)
            return None
        r = self.fetcher.get(url)
        if r is not None and r.status_code in (403, 503) and re.search(
                r"arvancloud|cloudflare|challenge|captcha|ddos", r.text[:5000], re.I):
            self.log.warning("blocked by firewall/challenge page (%s) at %s", r.status_code, url)
        return r

    def emit(self, rec: dict | None) -> bool:
        if rec and rec.get("name"):
            if self.writer.write(rec):
                if self.writer.count % 25 == 0:
                    self.log.info("%d products so far", self.writer.count)
                return True
        return False

    def load_robots(self):
        r = self.fetcher.get(urljoin(self.base, "/robots.txt"))
        if r is not None and r.status_code == 200 and "text/html" not in r.headers.get("content-type", ""):
            self.robots = Robots(r.text)
            if self.robots.crawl_delay and self.robots.crawl_delay > self.fetcher.delay:
                self.log.info("using crawl-delay %.1fs from robots.txt", self.robots.crawl_delay)
                self.fetcher.delay = self.robots.crawl_delay

    # ---- strategy 1: WooCommerce Store API
    def try_store_api(self) -> int:
        endpoints = [
            urljoin(self.base, "wp-json/wc/store/v1/products"),
            urljoin(self.base, "wp-json/wc/store/products"),
            urljoin(self.base, "?rest_route=/wc/store/v1/products"),
        ]
        for ep in endpoints:
            got, page, total_pages = 0, 1, None
            while True:
                sep = "&" if "?" in ep else "?"
                r = self.fetch(f"{ep}{sep}{urlencode({'per_page': 100, 'page': page})}")
                if r is None or r.status_code != 200:
                    break
                try:
                    data = r.json()
                except ValueError:
                    break
                if not isinstance(data, list) or not data:
                    break
                for item in data:
                    if isinstance(item, dict):
                        self.emit(product_from_store_api(item, self.host))
                got += len(data)
                if total_pages is None:
                    try:
                        total_pages = int(r.headers.get("X-WP-TotalPages") or 0)
                    except ValueError:
                        total_pages = 0
                    self.log.info("Store API works (%s pages)", total_pages or "?")
                if (total_pages and page >= total_pages) or len(data) < 100 or page >= 500:
                    break
                page += 1
            if got:
                return got
        return 0

    # ---- strategy 2: sitemaps
    @staticmethod
    def _is_product_sitemap(u: str) -> bool:
        name = unquote(urlparse(u).path.rsplit("/", 1)[-1]).lower()
        if "product" not in name and "محصول" not in name:
            return False
        return not re.search(r"product[_-]?(cat|category|tag|brand)|pa_|taxonom", name)

    @staticmethod
    def _is_useless_sitemap(u: str) -> bool:
        name = unquote(urlparse(u).path.rsplit("/", 1)[-1]).lower()
        return bool(re.search(r"author|user|attachment|tag|portfolio|post_tag|local-sitemap|video|image|news", name))

    def _read_sitemap(self, url: str):
        r = self.fetch(url)
        if r is None or r.status_code != 200:
            return None, []
        content = r.content
        if content[:2] == b"\x1f\x8b":
            try:
                content = gzip.decompress(content)
            except OSError:
                return None, []
        try:
            root = ET.fromstring(content)
        except ET.ParseError:
            return None, []
        kind = root.tag.rsplit("}", 1)[-1].lower()
        locs = [el.text.strip() for el in root.iter()
                if el.tag.rsplit("}", 1)[-1].lower() == "loc" and el.text and el.text.strip()]
        return kind, locs

    def sitemap_urls(self) -> tuple[list[str], bool]:
        """Returns (urls, urls_are_known_products)."""
        declared = list(self.robots.sitemaps)
        fallbacks = [urljoin(self.base, p) for p in (
            "sitemap_index.xml", "product-sitemap.xml", "wp-sitemap.xml", "sitemap.xml",
            "sitemap-products.xml", "sitemap/products.xml")]
        product_urls, other_urls, seen = set(), set(), set()

        def walk(start: str):
            q = deque([start])
            while q and len(seen) < 1000:
                sm = q.popleft()
                if sm in seen:
                    continue
                seen.add(sm)
                kind, locs = self._read_sitemap(sm)
                if kind is None:
                    continue
                if kind == "sitemapindex":
                    prod = [l for l in locs if self._is_product_sitemap(l)]
                    q.extend(prod or [l for l in locs if not self._is_useless_sitemap(l)])
                elif self._is_product_sitemap(sm):
                    product_urls.update(locs)
                else:
                    other_urls.update(locs)

        for sm in declared:
            walk(sm)
        for sm in fallbacks:
            if product_urls:
                break
            walk(sm)

        keep = lambda u: self.same_site(u) and not is_excluded(u)
        if product_urls:
            return sorted(u for u in product_urls if keep(u)), True
        others = [u for u in other_urls if keep(u)]
        guessed = [u for u in others if looks_like_product_url(u)]
        if guessed:
            return sorted(guessed), False
        return sorted(others), False

    # ---- page workers
    def _fetch_page(self, url: str):
        r = self.fetch(url)
        if r is None or r.status_code != 200:
            return url, None
        ctype = r.headers.get("content-type", "")
        if "html" not in ctype and ctype:
            return url, None
        r.encoding = r.encoding if r.encoding and r.encoding.lower() != "iso-8859-1" else "utf-8"
        return r.url, r.text

    def crawl_pages(self, urls: list[str], trust: bool) -> int:
        todo = [u for u in urls if u not in self.writer.seen_urls][: self.args.max_pages]
        self.log.info("parsing %d candidate pages from sitemap", len(todo))
        found = 0
        with ThreadPoolExecutor(max_workers=self.args.workers) as pool:
            futures = [pool.submit(self._fetch_page, u) for u in todo]
            for fut in as_completed(futures):
                final_url, html = fut.result()
                if html:
                    try:
                        rec = parse_product_page(html, final_url, self.host, trust_as_product=trust)
                    except Exception as e:  # never let one bad page kill the run
                        self.log.debug("parse error %s: %s", final_url, e)
                        rec = None
                    found += self.emit(rec)
        return found

    # ---- strategy 3: BFS link crawl
    def extract_links(self, html: str, page_url: str) -> list[str]:
        soup = BeautifulSoup(html, "lxml")
        out = []
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if href.startswith(("mailto:", "tel:", "javascript:", "#", "whatsapp:")):
                continue
            u = normalize_url(urljoin(page_url, href))
            if urlparse(u).scheme in ("http", "https") and self.same_site(u) and not is_excluded(u):
                out.append(u)
        return out

    def bfs(self) -> int:
        starts = [self.base] + [urljoin(self.base, p) for p in ("shop/", "products/", "store/", "فروشگاه/")]
        hi, lo = deque(), deque()
        seen = set()
        for s in starts:
            s = normalize_url(s)
            seen.add(s)
            hi.append(s)
        pages, found = 0, 0
        with ThreadPoolExecutor(max_workers=self.args.workers) as pool:
            while (hi or lo) and pages < self.args.max_pages and not self.stopped:
                batch = []
                while (hi or lo) and len(batch) < self.args.workers * 2:
                    batch.append(hi.popleft() if hi else lo.popleft())
                for final_url, html in pool.map(self._fetch_page, batch):
                    pages += 1
                    if not html:
                        continue
                    if final_url not in self.writer.seen_urls:
                        try:
                            found += self.emit(parse_product_page(html, final_url, self.host))
                        except Exception as e:
                            self.log.debug("parse error %s: %s", final_url, e)
                    for link in self.extract_links(html, final_url):
                        if link in seen:
                            continue
                        seen.add(link)
                        (hi if looks_like_product_url(link) or looks_like_listing_url(link) else lo).append(link)
                if pages % 100 < len(batch):
                    self.log.info("crawled %d pages, %d queued, %d products", pages, len(hi) + len(lo),
                                  self.writer.count)
        return found

    # ---- main
    def run(self) -> int:
        t0 = time.time()
        self.log.info("start")
        try:
            self.load_robots()
            home = self.fetcher.get(self.base)
            if home is None and self.base.startswith("https://") and not self.stopped:
                self.log.warning("https failed, trying http://")
                self.base = "http://" + self.base[len("https://"):]
                self.load_robots()
                home = self.fetcher.get(self.base)
            if home is None:
                self.log.error("site unreachable (network / DNS / geo-block?). Try --proxy.")
                return 0
            if home.status_code >= 400:
                self.log.warning("home page returned HTTP %s", home.status_code)

            mode, total = self.args.mode, 0
            if mode in ("auto", "api"):
                total = self.try_store_api()
                if total:
                    self.log.info("Store API: %d products", total)
            if not total and mode in ("auto", "sitemap") and not self.stopped:
                urls, trusted = self.sitemap_urls()
                if urls:
                    self.log.info("sitemap: %d urls (%s)", len(urls), "product sitemap" if trusted else "guessed")
                    total = self.crawl_pages(urls, trusted)
            if not total and mode in ("auto", "crawl") and not self.stopped:
                self.log.info("falling back to link crawling (max %d pages)", self.args.max_pages)
                total = self.bfs()
            if self.blocked_by_robots:
                self.log.warning("%d URLs skipped because robots.txt disallows them "
                                 "(use --ignore-robots only if you have the owner's permission)",
                                 self.blocked_by_robots)
            if self.stopped:
                self.log.warning("stopped by user")
            self.log.info("done: %d products in %.0fs", self.writer.count, time.time() - t0)
            return self.writer.count
        finally:
            self.writer.close()


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser(description="Crawl products from WooCommerce / generic shop sites.")
    ap.add_argument("--sites", nargs="*", help="site URLs (default: the built-in list)")
    ap.add_argument("--sites-file", help="text file with one site URL per line")
    ap.add_argument("--out", default="output", help="output folder (default: output)")
    ap.add_argument("--mode", choices=["auto", "api", "sitemap", "crawl"], default="auto")
    ap.add_argument("--delay", type=float, default=1.0, help="seconds between requests per site (default 1.0)")
    ap.add_argument("--workers", type=int, default=3, help="parallel requests per site (default 3)")
    ap.add_argument("--parallel-sites", type=int, default=4, help="sites crawled at the same time (default 4)")
    ap.add_argument("--max-pages", type=int, default=5000, help="max HTML pages per site (default 5000)")
    ap.add_argument("--timeout", type=float, default=30)
    ap.add_argument("--proxy", help="http(s)/socks proxy, e.g. http://127.0.0.1:8080")
    ap.add_argument("--user-agent", default=DEFAULT_USER_AGENT)
    ap.add_argument("--resume", action="store_true", help="keep existing .jsonl and skip already-saved URLs")
    ap.add_argument("--ignore-robots", action="store_true",
                    help="do not obey robots.txt (only with the site owner's permission)")
    ap.add_argument("--insecure", action="store_true", help="skip TLS certificate verification")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(name)-22s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")
    logging.getLogger("urllib3").setLevel(logging.WARNING)

    sites = list(args.sites or [])
    if args.sites_file:
        with open(args.sites_file, encoding="utf-8") as f:
            sites += [l.strip() for l in f if l.strip() and not l.startswith("#")]
    if not sites:
        sites = DEFAULT_SITES
    sites = _dedupe(s if s.startswith("http") else "https://" + s for s in sites)

    os.makedirs(args.out, exist_ok=True)
    crawlers = [SiteCrawler(s, args) for s in sites]
    summary = {}
    with ThreadPoolExecutor(max_workers=max(1, args.parallel_sites)) as pool:
        futs = {pool.submit(c.run): c for c in crawlers}
        for fut in as_completed(futs):
            c = futs[fut]
            try:
                summary[c.host] = fut.result()
            except Exception as e:
                logging.getLogger(c.host).exception("crashed: %s", e)
                summary[c.host] = 0

    all_rows = []
    for c in crawlers:
        rows = read_jsonl(c.writer.path)
        if rows:
            write_csv(rows, os.path.join(args.out, f"{c.host}.csv"))
        all_rows += rows
    write_csv(all_rows, os.path.join(args.out, "all_products.csv"))
    xlsx = write_xlsx(all_rows, os.path.join(args.out, "all_products.xlsx"))

    print("\n===== summary =====")
    for c in crawlers:
        print(f"  {c.host:<28} {summary.get(c.host, 0):>6} products")
    print(f"  {'TOTAL':<28} {len(all_rows):>6}")
    print(f"\nfiles in: {os.path.abspath(args.out)}" + ("" if xlsx else "  (pip install openpyxl for .xlsx)"))


if __name__ == "__main__":
    main()
